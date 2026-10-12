"""apps.common.excel 单元测试（不依赖业务 model，可独立运行）。

覆盖：写读往返、表头别名、csv utf-8-sig/gbk 回退、空行跳过、缺必填列报错、
多表簿构建、列宽、下载响应。
"""
import io

import pytest
from openpyxl import load_workbook

from apps.common.excel import (
    ExcelParseError,
    build_export_response,
    build_workbook,
    read_raw_rows,
    read_rows,
    set_column_widths,
    workbook_to_bytes,
    write_sheet,
)


def _save(wb) -> bytes:
    return workbook_to_bytes(wb)


def test_write_and_read_roundtrip():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = 't'
    write_sheet(ws, ['name', 'age'], [['张三', '30'], ['李四', '25']])
    raw = _save(wb)

    rows = read_rows(raw)
    assert rows == [
        {'name': '张三', 'age': '30'},
        {'name': '李四', 'age': '25'},
    ]


def test_header_aliases_and_columns():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_sheet(ws, ['姓名', '年龄'], [['张三', '30']])
    raw = _save(wb)

    rows = read_rows(
        raw,
        aliases={'name': ('姓名', 'name'), 'age': ('年龄', 'age')},
        columns=['name', 'age'],
    )
    assert rows == [{'name': '张三', 'age': '30'}]
    assert set(rows[0].keys()) == {'name', 'age'}


def test_csv_utf8sig_fallback():
    buf = io.StringIO()
    import csv

    w = csv.writer(buf)
    w.writerow(['name', 'city'])
    w.writerow(['王五', '北京'])
    raw = buf.getvalue().encode('utf-8-sig')

    rows = read_rows(raw)
    assert rows == [{'name': '王五', 'city': '北京'}]


def test_csv_gbk_fallback():
    buf = io.StringIO()
    import csv

    w = csv.writer(buf)
    w.writerow(['name', 'city'])
    w.writerow(['赵六', '上海'])
    raw = buf.getvalue().encode('gbk')

    rows = read_rows(raw)
    assert rows == [{'name': '赵六', 'city': '上海'}]


def test_skip_blank_rows():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_sheet(ws, ['name'], [['a'], [''], ['c']])
    raw = _save(wb)

    rows = read_rows(raw)
    assert [r['name'] for r in rows] == ['a', 'c']


def test_missing_required_raises():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_sheet(ws, ['name'], [['a']])
    raw = _save(wb)

    with pytest.raises(ExcelParseError):
        read_rows(raw, required=['name', 'missing_col'])


def test_build_workbook_multi_sheet():
    wb = build_workbook({
        '用户': (['name'], [['u1']]),
        '部门': (['dept'], [['d1']]),
    })
    raw = _save(wb)
    reopened = load_workbook(io.BytesIO(raw))
    assert reopened.sheetnames == ['用户', '部门']
    assert reopened['用户']['A2'].value == 'u1'


def test_set_column_widths():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    set_column_widths(ws, [10, 20, 5])
    assert ws.column_dimensions['A'].width == 10
    assert ws.column_dimensions['C'].width == 5


def test_workbook_to_bytes_nonempty():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_sheet(ws, ['x'], [['1']])
    assert isinstance(_save(wb), bytes) and len(_save(wb)) > 0


def test_build_export_response():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_sheet(ws, ['x'], [['1']])
    resp = build_export_response(wb, 'demo.xlsx')
    assert resp['Content-Type'].startswith('application/vnd.openxmlformats')
    assert 'attachment' in resp['Content-Disposition']
    assert 'demo.xlsx' in resp['Content-Disposition']
    assert len(resp.content) > 0


def test_read_raw_rows_keeps_header_and_returns_tuples():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_sheet(ws, ['name', 'age'], [['张三', '30'], ['李四', '25']])
    raw = _save(wb)

    rows = read_raw_rows(raw)
    assert isinstance(rows, list)
    assert all(isinstance(r, tuple) for r in rows)
    assert rows[0] == ('name', 'age')        # 表头行保留（区别于 read_rows 的 dict 视图）
    assert rows[1] == ('张三', '30')
    assert rows[2] == ('李四', '25')
    assert len(rows) == 3


def test_read_raw_rows_equivalence_to_bare_openpyxl():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_sheet(ws, ['name', 'age'], [['张三', '30']])
    raw = _save(wb)

    expected = [
        tuple(r) for r in load_workbook(
            io.BytesIO(raw), read_only=True, data_only=True,
        ).active.iter_rows(values_only=True)
    ]
    # 与 import_schools 旧写法的 load_workbook(read_only, data_only) 逐字节等价
    assert read_raw_rows(raw) == expected


def test_read_raw_rows_from_file_path(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_sheet(ws, ['a', 'b'], [['1', '2']])
    p = tmp_path / 'sample.xlsx'
    p.write_bytes(_save(wb))

    rows = read_raw_rows(str(p))
    assert rows[0] == ('a', 'b')
    assert rows[1] == ('1', '2')


def test_read_raw_rows_csv_keeps_header():
    buf = io.StringIO()
    import csv

    w = csv.writer(buf)
    w.writerow(['name', 'city'])
    w.writerow(['王五', '北京'])
    raw = buf.getvalue().encode('utf-8-sig')

    rows = read_raw_rows(raw)
    assert rows[0] == ('name', 'city')
    assert rows[1] == ('王五', '北京')
