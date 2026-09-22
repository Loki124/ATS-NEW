"""原因标签 导入 / 模板 xlsx+csv 双格式单测。

覆盖:
- xlsx 导入成功 (含中文与特殊字符)
- csv 导入成功 (回归, 含 utf-8-sig BOM)
- 缺 name 列 → 40002 CSV_FORMAT_INVALID (xlsx / csv 两条路径)
- 空白行跳过 (Excel 尾部空行不计为错误)
- 模板下载: 默认 xlsx / ?format=csv (保留 BOM)
- gbk 编码 csv 仍可用

 fixture 复用根目录 conftest (admin_api_client = (client, user))。
"""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from apps.reason_library.models import ReasonTag

pytestmark = pytest.mark.django_db

TAG_IMPORT = '/api/v1/reason-library/tags/import/'
TAG_TEMPLATE = '/api/v1/reason-library/tags/import-template/'
HEADERS = ['name', 'en_name', 'tip', 'type', 'enabled']


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def make_csv(content: str, name: str = 'tags.csv', encoding: str = 'utf-8-sig'):
    """构造 csv 上传对象。"""
    upload = io.BytesIO(content.encode(encoding))
    upload.name = name
    return upload


def make_xlsx(rows, name: str = 'tags.xlsx', headers=None, sheet_name='原因标签'):
    """构造 xlsx 上传对象。rows: list[list] 数据行 (表头下方)。"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers if headers is not None else HEADERS)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    upload = io.BytesIO(buf.getvalue())
    upload.name = name
    return upload


# --------------------------------------------------------------------------- #
# xlsx import
# --------------------------------------------------------------------------- #
def test_xlsx_import_creates_tags(admin_api_client):
    """Happy path: xlsx 导入创建标签, 字段完整落库。"""
    client, _ = admin_api_client
    upload = make_xlsx([
        ['xlsx-tag-1', 'Xlsx Tag 1', '提示1', 'custom', 'true'],
        ['xlsx-tag-2', 'Xlsx Tag 2', '', 'custom', 'false'],
    ])
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 200, resp.content
    body = resp.json()
    assert body['code'] == 0
    assert body['data']['created'] == 2
    assert body['data']['errors'] == []

    t1 = ReasonTag.objects.get(name='xlsx-tag-1')
    assert t1.en_name == 'Xlsx Tag 1'
    assert t1.tip == '提示1'
    assert t1.enabled is True
    t2 = ReasonTag.objects.get(name='xlsx-tag-2')
    assert t2.enabled is False
    assert t2.tip == ''


def test_xlsx_import_chinese_and_special_chars(admin_api_client):
    """中文 + 逗号/引号/emoji 等特殊字符不被破坏。"""
    client, _ = admin_api_client
    weird = '面试取消, "候选人" <离职>\U0001F600'
    upload = make_xlsx([[weird, 'Special, "quoted"', '含,逗号 的提示', 'custom', 'true']])
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 1
    tag = ReasonTag.objects.get(name=weird)
    assert tag.en_name == 'Special, "quoted"'
    assert tag.tip == '含,逗号 的提示'


def test_xlsx_import_skips_blank_rows(admin_api_client):
    """Excel 常见的尾部/中间空行 → 跳过, 不计为错误。"""
    client, _ = admin_api_client
    upload = make_xlsx([
        ['xlsx-blank-1', '', '', 'custom', 'true'],
        ['', '', '', '', ''],
        [None, None, None, None, None],
        ['xlsx-blank-2', '', '', 'custom', 'true'],
    ])
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 200, resp.content
    body = resp.json()
    assert body['data']['created'] == 2
    assert body['data']['errors'] == []


def test_xlsx_import_missing_name_column_returns_40002(admin_api_client):
    """缺 name 列 → 400 + 40002 CSV_FORMAT_INVALID。"""
    client, _ = admin_api_client
    upload = make_xlsx([['x']], headers=['en_name', 'tip'])
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 400
    body = resp.json()
    assert body['code'] == 40002
    assert 'name' in body['message']


def test_xlsx_import_broken_file_returns_40002(admin_api_client):
    """非 Excel 内容却用 .xlsx 扩展名 → 40002 (不是 500)。"""
    client, _ = admin_api_client
    upload = io.BytesIO(b'this-is-not-a-real-xlsx')
    upload.name = 'broken.xlsx'
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 400
    assert resp.json()['code'] == 40002


def test_xlsx_import_duplicate_returns_40001(admin_api_client, reason_tag):
    """同名 (库内已有) → 40001 TAG_NAME_DUPLICATED。"""
    client, _ = admin_api_client
    upload = make_xlsx([[reason_tag.name, 'dup', '', 'custom', 'true']])
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 400
    assert resp.json()['code'] == 40001


def test_xlsx_import_enabled_accepts_numeric(admin_api_client):
    """Excel 里 enabled 填数字/布尔 → 字符串化后仍能正确解析。"""
    client, _ = admin_api_client
    upload = make_xlsx([
        ['xlsx-num-1', '', '', 'custom', 1],
        ['xlsx-num-2', '', '', 'custom', 0],
    ])
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 2
    assert ReasonTag.objects.get(name='xlsx-num-1').enabled is True
    assert ReasonTag.objects.get(name='xlsx-num-2').enabled is False


# --------------------------------------------------------------------------- #
# csv import (回归)
# --------------------------------------------------------------------------- #
def test_csv_import_still_works(admin_api_client):
    client, _ = admin_api_client
    csv_content = 'name,en_name,tip,type,enabled\ncsv-tag-x1,CSV X1,tip1,custom,true\n'
    upload = make_csv(csv_content)
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 1


def test_csv_import_bom_and_gbk(admin_api_client):
    """带 BOM 的 utf-8 与 gbk 编码 csv 都能被正确解码。"""
    client, _ = admin_api_client
    upload = make_csv(
        'name,en_name,tip,type,enabled\ncsv-bom-tag,BOM Tag,带 BOM 的提示,custom,true\n',
        name='bom.csv',
    )
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert ReasonTag.objects.get(name='csv-bom-tag').tip == '带 BOM 的提示'

    client2, _ = admin_api_client
    upload = make_csv(
        'name,en_name,tip,type,enabled\ngbk-tag,GBK Tag,中文提示,custom,false\n',
        name='gbk.csv', encoding='gbk',
    )
    resp = client2.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert ReasonTag.objects.get(name='gbk-tag').enabled is False


def test_csv_import_missing_name_column_returns_40002(admin_api_client):
    client, _ = admin_api_client
    upload = make_csv('en_name,tip\nX,x\n', name='bad.csv')
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 400
    assert resp.json()['code'] == 40002


# --------------------------------------------------------------------------- #
# 模板下载
# --------------------------------------------------------------------------- #
def test_template_default_is_xlsx(admin_api_client):
    client, _ = admin_api_client
    resp = client.get(TAG_TEMPLATE)
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp['Content-Type']
    assert 'reason-tags-import-template.xlsx' in resp['Content-Disposition']
    # xlsx = zip magic "PK"
    assert resp.content[:2] == b'PK'


def test_template_xlsx_can_be_parsed_back(admin_api_client):
    """下载的 xlsx 模板回填数据后应能直接被导入解析 (闭环)。"""
    client, _ = admin_api_client
    resp = client.get(TAG_TEMPLATE)
    wb = None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][:5] == tuple(HEADERS)
        assert rows[1][0] == '示例标签'
    finally:
        if wb is not None:
            wb.close()


def test_template_csv_keeps_bom(admin_api_client):
    client, _ = admin_api_client
    resp = client.get(TAG_TEMPLATE, {'format': 'csv'})
    assert resp.status_code == 200
    assert 'reason-tags-import-template.csv' in resp['Content-Disposition']
    text = resp.content.decode('utf-8')
    assert text.startswith('\ufeff')  # BOM 必须保留 (Excel 兼容铁律)
    lines = text.lstrip('\ufeff').strip().splitlines()
    assert lines[0] == 'name,en_name,tip,type,enabled'
    assert '示例标签' in lines[1]


def test_template_csv_is_importable(admin_api_client):
    """csv 模板回填数据 → 导入成功。"""
    client, _ = admin_api_client
    resp = client.get(TAG_TEMPLATE, {'format': 'csv'})
    text = resp.content.decode('utf-8-sig')
    header_line = text.strip().splitlines()[0]
    csv_content = f'{header_line}\nfrom-template,From Template,,custom,true\n'
    upload = make_csv(csv_content, name='from-template.csv')
    resp = client.post(TAG_IMPORT, {'file': upload}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert ReasonTag.objects.filter(name='from-template').exists()
