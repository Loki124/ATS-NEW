"""QA 独立复验探针（临时文件，用完即删）。

不复用工程师的 test_tag_io.py，全部自行构造，重点打边界：
sheet 名 / 列乱序 / 多余列 / 大小写 / CRLF / 空文件 / 扩展名欺骗 / 超长 name /
响应契约字段集合 / 错误码。
"""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from apps.reason_library.models import ReasonTag

pytestmark = pytest.mark.django_db

IMPORT = '/api/v1/reason-library/tags/import/'
TEMPLATE = '/api/v1/reason-library/tags/import-template/'


def mk_xlsx(rows, name='tags.xlsx', headers=None, sheet_title='原因标签', extra_sheet=False):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers if headers is not None else ['name', 'en_name', 'tip', 'type', 'enabled'])
    for r in rows:
        ws.append(r)
    if extra_sheet:
        other = wb.create_sheet('别的表')
        other.append(['x', 'y'])
        other.append([1, 2])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    up = io.BytesIO(buf.getvalue())
    up.name = name
    return up


def mk_csv(content, name='t.csv', encoding='utf-8-sig'):
    up = io.BytesIO(content.encode(encoding))
    up.name = name
    return up


# ---------- A. xlsx 路径 ----------
def test_qa_xlsx_default_sheet_title(admin_api_client):
    """没有『原因标签』工作表（用户用 openpyxl/其他工具另存后表名为 Sheet1）仍应能导。"""
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [['qa-sheet1-a', '', '', 'custom', 'true']], sheet_title='Sheet1')}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 1


def test_qa_xlsx_multi_sheet_active_is_data(admin_api_client):
    """多工作表且数据在第一张——不应被其他表干扰。"""
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [['qa-multi-a', '', '', 'custom', 'true']], extra_sheet=True)}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 1


def test_qa_xlsx_reordered_and_extra_columns(admin_api_client):
    """列顺序打乱 + 多一列未知列 → 按表头映射，未知列忽略。"""
    client, _ = admin_api_client
    headers = ['tip', 'unknown_col', 'name', 'enabled', 'en_name', 'type']
    rows = [['提示-X', 'zzz', 'qa-reorder-a', 'false', 'Reorder A', 'custom']]
    resp = client.post(IMPORT, {'file': mk_xlsx(rows, headers=headers)}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 1
    t = ReasonTag.objects.get(name='qa-reorder-a')
    assert t.en_name == 'Reorder A'
    assert t.tip == '提示-X'
    assert t.enabled is False


def test_qa_xlsx_uppercase_extension(admin_api_client):
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [['qa-upper-a', '', '', 'custom', 'true']], name='TAGS.XLSX')}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 1


def test_qa_xlsm_extension(admin_api_client):
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [['qa-xlsm-a', '', '', 'custom', 'true']], name='tags.xlsm')}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 1


def test_qa_xlsx_only_header_no_rows(admin_api_client):
    """只有表头、零数据行 → 200，created=0，errors 空。"""
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx([])}, format='multipart')
    assert resp.status_code == 200, resp.content
    body = resp.json()
    assert body['data']['created'] == 0
    assert body['data']['errors'] == []


def test_qa_xlsx_empty_file_40002(admin_api_client):
    """0 字节 .xlsx → 40002，不得 500。"""
    client, _ = admin_api_client
    up = io.BytesIO(b'')
    up.name = 'empty.xlsx'
    resp = client.post(IMPORT, {'file': up}, format='multipart')
    assert resp.status_code == 400
    assert resp.json()['code'] == 40002


def test_qa_xlsx_missing_name_header_with_alias(admin_api_client):
    """表头写成『标签名称』之类的中文别名 → 按契约判 40002 并提示缺列。"""
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [['x']], headers=['标签名称', 'en_name'])}, format='multipart')
    assert resp.status_code == 400
    body = resp.json()
    assert body['code'] == 40002
    assert 'name' in body['message']


def test_qa_real_xlsx_named_csv_not_500(admin_api_client):
    """真 xlsx 二进制却后缀 .csv → 必须 400/40002，不能 500。"""
    client, _ = admin_api_client
    up = mk_xlsx([['qa-x', '', '', 'custom', 'true']], name='confuse.csv')
    resp = client.post(IMPORT, {'file': up}, format='multipart')
    assert resp.status_code == 400, resp.content
    assert resp.json()['code'] == 40002


# ---------- B. csv 路径（回归） ----------
def test_qa_csv_crlf_and_trailing_blank_lines(admin_api_client):
    """CRLF + 单个 BOM + 尾部多个空行不应产生错误行。"""
    client, _ = admin_api_client
    content = (
        '\ufeffname,en_name,tip,type,enabled\r\n'
        'qa-crlf-a,CRLF A,tip-a,custom,true\r\n'
        '\r\n'
        'qa-crlf-b,CRLF B,,custom,false\r\n'
        '\r\n\r\n'
    )
    resp = client.post(IMPORT, {'file': mk_csv(content, encoding='utf-8')}, format='multipart')
    assert resp.status_code == 200, resp.content
    body = resp.json()
    assert body['data']['created'] == 2, body
    assert body['data']['errors'] == []
    assert ReasonTag.objects.get(name='qa-crlf-b').enabled is False


def test_qa_csv_double_bom_is_rejected_preexisting(admin_api_client):
    """双 BOM（手工 \ufeff + utf-8-sig 再编码一层）→ 40002。

    注意: 改动前 `upload.read().decode('utf-8-sig')` 只剥一层 BOM, 行为完全一致,
    故这是**既有行为, 不是本次改造引入的回归**。此用例仅固化事实。
    """
    client, _ = admin_api_client
    content = '\ufeffname,en_name,tip,type,enabled\nqa-dblbom,,,custom,true\n'
    resp = client.post(IMPORT, {'file': mk_csv(content, encoding='utf-8-sig')}, format='multipart')
    print(f'\n[QA-PROBE] double-bom status={resp.status_code}')
    assert resp.status_code == 400
    assert resp.json()['code'] == 40002


def test_qa_csv_quoted_comma_and_newline_field(admin_api_client):
    """CSV 引号包裹的逗号（由 csv.writer 正确转义内层引号）→ 字段不被切碎。"""
    import csv as _csv
    client, _ = admin_api_client
    weird = 'qa-逗号, "引号" \U0001F3AF'
    buf = io.StringIO()
    w = _csv.writer(buf)
    w.writerow(['name', 'en_name', 'tip', 'type', 'enabled'])
    w.writerow([weird, 'EN, quoted', 'tip, x', 'custom', 'true'])
    resp = client.post(IMPORT, {'file': mk_csv(buf.getvalue())}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 1
    tag = ReasonTag.objects.get(name=weird)
    assert tag.en_name == 'EN, quoted'
    assert tag.tip == 'tip, x'


def test_qa_csv_unescaped_inner_quote_loosely_parsed(admin_api_client):
    """内层引号未转义（用户手工拼的脏 CSV）→ Python csv 宽松吸收，不当成崩溃/500。

    仅固化事实: 名字会被解析成不含外侧引号的形式, 创建成功, 不 500。
    """
    client, _ = admin_api_client
    weird = 'qa-逗号, "引号" \U0001F3AF'
    content = f'name,en_name,tip,type,enabled\n"{weird}","EN, quoted","tip, x",custom,true\n'
    resp = client.post(IMPORT, {'file': mk_csv(content)}, format='multipart')
    print(f'\n[QA-PROBE] loose-quote status={resp.status_code}')
    assert resp.status_code == 200, resp.content
    created = list(ReasonTag.objects.filter(name__startswith='qa-逗号').values_list('name', flat=True))
    print(f'[QA-PROBE] loose-quote stored name={created!r}')
    assert created, '脏 CSV 应被宽松解析而非 500'


def test_qa_csv_gbk_still_ok(admin_api_client):
    client, _ = admin_api_client
    content = 'name,en_name,tip,type,enabled\nqa-gbk-a,GBK,中文提示,custom,false\n'
    resp = client.post(IMPORT, {'file': mk_csv(content, encoding='gbk')}, format='multipart')
    assert resp.status_code == 200, resp.content
    tag = ReasonTag.objects.get(name='qa-gbk-a')
    assert tag.tip == '中文提示'
    assert tag.enabled is False


def test_qa_csv_header_only(admin_api_client):
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_csv('name,en_name,tip,type,enabled\n')}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == 0


def test_qa_csv_empty_bytes_40002(admin_api_client):
    client, _ = admin_api_client
    up = io.BytesIO(b'')
    up.name = 'empty.csv'
    resp = client.post(IMPORT, {'file': up}, format='multipart')
    assert resp.status_code == 400
    assert resp.json()['code'] == 40002


def test_qa_file_with_no_name_falls_back_to_csv(admin_api_client):
    """UploadedFile.name 缺失 → 走 csv 分支，不 500。"""
    client, _ = admin_api_client
    up = io.BytesIO('name\naaa\n'.encode('utf-8'))
    up.name = ''
    resp = client.post(IMPORT, {'file': up}, format='multipart')
    assert resp.status_code in (200, 400), resp.content


# ---------- C. 校验 / 错误码契约 ----------
def test_qa_type_system_row_40002(admin_api_client):
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [['qa-system-row', '', '', 'system', 'true']])}, format='multipart')
    assert resp.status_code == 400
    assert resp.json()['code'] == 40002


def test_qa_duplicate_inside_file_partial_success(admin_api_client):
    """同文件内重复 → created=1 + errors 1 条，HTTP 200（部分成功不算全失败）。"""
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx([
        ['qa-dup-same', '', '', 'custom', 'true'],
        ['qa-dup-same', '', '', 'custom', 'true'],
    ])}, format='multipart')
    assert resp.status_code == 200, resp.content
    body = resp.json()
    assert body['data']['created'] == 1
    assert len(body['data']['errors']) == 1


def test_qa_response_contract_keys_and_message(admin_api_client):
    """对外契约：data 恰为 {created, skipped, errors}，message = 成功导入 N 条。"""
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx([
        ['qa-contract-a', '', '', 'custom', 'true'],
        ['qa-contract-b', '', '', 'custom', 'true'],
    ])}, format='multipart')
    assert resp.status_code == 200, resp.content
    body = resp.json()
    assert set(body['data'].keys()) == {'created', 'skipped', 'errors'}
    assert body['data']['created'] == 2
    assert body['message'] == '成功导入 2 条'


def test_qa_missing_file_field_40002(admin_api_client):
    """字段名不是 file → 40002『未上传文件』。"""
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'wrong': mk_xlsx([['x', '', '', 'custom', 'true']])}, format='multipart')
    assert resp.status_code == 400
    assert resp.json()['code'] == 40002


def test_qa_enabled_parsing_matrix(admin_api_client):
    """enabled 多形态解析。"""
    client, _ = admin_api_client
    cases = [
        ('qa-en-1', 'TRUE', True), ('qa-en-2', 'Yes', True),
        ('qa-en-3', ' no', False), ('qa-en-4', '0', False),
        ('qa-en-5', '', True), ('qa-en-6', True, True),
        ('qa-en-7', 1, True), ('qa-en-8', 0, False),
    ]
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [[n, '', '', 'custom', v] for n, v, _ in cases])}, format='multipart')
    assert resp.status_code == 200, resp.content
    assert resp.json()['data']['created'] == len(cases)
    for n, _, expected in cases:
        assert ReasonTag.objects.get(name=n).enabled is expected, n


# ---------- D. 模板下载契约 ----------
def test_qa_template_default_xlsx_magic_and_filename(admin_api_client):
    client, _ = admin_api_client
    resp = client.get(TEMPLATE)
    assert resp.status_code == 200
    assert resp['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'filename="reason-tags-import-template.xlsx"' in resp['Content-Disposition']
    assert resp.content[:2] == b'PK'


def test_qa_template_xlsx_has_two_sheets_and_color(admin_api_client):
    client, _ = admin_api_client
    resp = client.get(TEMPLATE)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    try:
        assert '原因标签' in wb.sheetnames
        assert '填写说明' in wb.sheetnames
        ws = wb['原因标签']
        assert [c.value for c in ws[1]] == ['name', 'en_name', 'tip', 'type', 'enabled']
        assert ws['A1'].fill.start_color.rgb.endswith('6366F1')
    finally:
        wb.close()


def test_qa_template_csv_equal_columns(admin_api_client):
    client, _ = admin_api_client
    resp = client.get(TEMPLATE, {'format': 'csv'})
    assert resp.status_code == 200
    text = resp.content.decode('utf-8')
    assert text.startswith('\ufeff')
    assert 'filename="reason-tags-import-template.csv"' in resp['Content-Disposition']
    assert text.lstrip('\ufeff').splitlines()[0] == 'name,en_name,tip,type,enabled'


def test_qa_template_unknown_format_falls_back_to_xlsx(admin_api_client):
    """非法 format 参数 → 回退 xlsx（不 500）。"""
    client, _ = admin_api_client
    resp = client.get(TEMPLATE, {'format': 'pdf'})
    assert resp.status_code == 200
    assert resp.content[:2] == b'PK'


# ---------- E. 超长 name（探针，观察是否会 500） ----------
def test_qa_name_exceeding_max_length(admin_api_client):
    """name 超 32 字符：观察是否 500（DataError 未被 IntegrityError 捕获）。"""
    client, _ = admin_api_client
    long_name = '长' * 60
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [[long_name, '', '', 'custom', 'true']])}, format='multipart')
    print(f'\n[QA-PROBE] long-name status={resp.status_code} body={resp.content[:300]!r}')
    assert resp.status_code in (200, 400), f'疑似 500: {resp.status_code} {resp.content[:300]!r}'


def test_qa_en_name_and_tip_truncation(admin_api_client):
    """en_name >64 / tip >128 应被截断而非报错。"""
    client, _ = admin_api_client
    resp = client.post(IMPORT, {'file': mk_xlsx(
        [['qa-trunc', 'E' * 100, 'T' * 200, 'custom', 'true']])}, format='multipart')
    assert resp.status_code == 200, resp.content
    t = ReasonTag.objects.get(name='qa-trunc')
    assert len(t.en_name) == 64
    assert len(t.tip) == 128
