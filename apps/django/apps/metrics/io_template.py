"""指标模板导入/导出工具（xlsx + csv 双格式，零额外重依赖）。

照搬 campus_control/io_indicator.py 的设计：模板与导入共用同一套列结构，保证
「下载模板 → 填数据 → 导入」闭环一致。

列结构（行 = 一条指标模板）：
  0 模板名称        (template name，唯一)
  1 指标类型        (原子 / 派生；留空则先查原子再查派生，同名优先原子)
  2 引用指标名称    (atomic / derived metric name，须已存在)
  3 支持的运算符     (逗号分隔的英文 code 或中文，如 IS_EMPTY,大于)
  4 参数枚举        (逗号分隔的枚举值，留空表示非枚举)
  5 允许为空        (是/否，留空默认 否)
  6 状态            (启用/停用，留空默认 启用)
  7 说明            (自由文本)

注意：参数配置(param_config) / 值域配置(value_domain) 是自由 JSON，导入暂不强制，
留空即可；导出也只导这 8 列的业务字段，减轻复杂度（后续如需再加）。
"""
import base64
import csv
import io
import logging

from django.db import DatabaseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from apps.audit.models import AuditLog
from apps.rule_engine.models import UnifiedOperator

from .models import AtomicMetric, DerivedMetric, MetricTemplate

logger = logging.getLogger(__name__)


# ---- 列头（模板首行 / 导出首行 一致） ----
TEMPLATE_HEADERS = ['模板名称', '指标类型', '引用指标名称', '支持的运算符', '参数枚举', '允许为空', '状态', '说明']

# 样式（与 io_indicator 对齐视觉）
HEADER_FILL = PatternFill('solid', fgColor='6366F1')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=11)
SAMPLE_FILL = PatternFill('solid', fgColor='F1F5F9')
THIN = Side(style='thin', color='E2E8F0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BOOL_TRUE = ('是', '启用', 'true', 'True', 'TRUE', '1', 'yes', 'Yes', 'YES', 'y', 'Y')
BOOL_FALSE = ('否', '停用', 'false', 'False', 'FALSE', '0', 'no', 'No', 'NO', 'n', 'N')

VALID_TYPES = ('原子', '派生')

# 中文算子 → code 映射（解析运算符用，也接受英文 code 直传）
OPERATOR_CN_MAP = {
    '': 'IS_EMPTY',
    '空': 'IS_EMPTY', '为空': 'IS_EMPTY', '空值': 'IS_EMPTY',
    '非空': 'IS_NOT_EMPTY', '不为空': 'IS_NOT_EMPTY',
    '包含': 'IN', '属于': 'IN', '属于集合': 'IN',
    '不包含': 'NOT_IN', '不属于': 'NOT_IN', '不属于集合': 'NOT_IN',
    '大于': 'GT', '大于等于': 'GTE', '等于': 'EQ', '不等于': 'NEQ',
    '小于': 'LT', '小于等于': 'LTE', '介于': 'BETWEEN', '区间': 'BETWEEN',
}

ALL_OPERATORS = set(UnifiedOperator.values)


# ===========================================================================
#  布尔 / 状态解析
# ===========================================================================
def _parse_bool(raw):
    """解析「允许为空」为 bool。空 → False（默认不允许空）；非法 → None。"""
    if raw is None:
        return False
    v = str(raw).strip()
    if v == '':
        return False
    if v in BOOL_TRUE:
        return True
    if v in BOOL_FALSE:
        return False
    return None


def _bool_to_text(v):
    return '是' if v else '否'


def _parse_status(raw):
    """解析「状态」为 'enabled' / 'disabled' / None(非法)。空 → enabled。"""
    if raw is None:
        return 'enabled'
    v = str(raw).strip()
    if v == '':
        return 'enabled'
    if v in ('启用', '是', 'true', 'True', 'TRUE', '1', 'yes', 'Yes', 'YES'):
        return 'enabled'
    if v in ('停用', '否', 'false', 'False', 'FALSE', '0', 'no', 'No', 'NO'):
        return 'disabled'
    return None


def _status_to_text(status):
    return '启用' if status == 'enabled' else '停用'


def _normalize_operators(raw):
    """逗号分隔的中/英算子字符串 → code 列表；非法 code → (None, error_msg)。"""
    if raw is None:
        return [], None
    text = str(raw).strip()
    if text == '':
        return [], None
    codes = []
    for part in text.split(','):
        p = part.strip()
        if p == '':
            continue
        if p in ALL_OPERATORS:
            codes.append(p)
            continue
        code = OPERATOR_CN_MAP.get(p)
        if code is None:
            return None, f'不支持的运算符：{p}'
        codes.append(code)
    if not codes:
        return None, '未提供任何合法运算符'
    return codes, None


# ===========================================================================
#  行构建（导出用）
# ===========================================================================
def _template_row(t):
    """MetricTemplate → 导出行 dict。"""
    metric = t.metric
    metric_name = metric.name if metric else ''
    metric_kind = '原子' if t.metric_kind == 'atomic' else '派生'
    return {
        'name': t.name,
        'metric_kind': metric_kind,
        'metric_name': metric_name,
        'operators': ','.join(t.operators or []),
        'param_enums': ','.join(t.param_enums or []),
        'param_allow_null': _bool_to_text(t.param_allow_null),
        'status': _status_to_text(t.status),
        'description': t.description or '',
    }


def build_template_export_rows():
    """从库内取全部模板，按 name 排序，返回导出行 list。"""
    return [
        _template_row(t)
        for t in MetricTemplate.objects.select_related('atomic_metric', 'derived_metric').all().order_by('name')
    ]


# ===========================================================================
#  xlsx 构建
# ===========================================================================
def _write_header(ws):
    for c, h in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal='center')


def _style_sheet(ws):
    for c, w in enumerate([20, 12, 20, 36, 20, 12, 12, 40], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


def build_template_export_workbook(rows):
    """rows: list[{name, metric_kind, metric_name, operators, param_enums, param_allow_null, status, description}] → xlsx 工作簿。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '指标模板'
    _write_header(ws)
    for i, r in enumerate(rows):
        row_idx = i + 2
        ws.cell(row=row_idx, column=1, value=r.get('name', ''))
        ws.cell(row=row_idx, column=2, value=r.get('metric_kind', ''))
        ws.cell(row=row_idx, column=3, value=r.get('metric_name', ''))
        ws.cell(row=row_idx, column=4, value=r.get('operators', ''))
        ws.cell(row=row_idx, column=5, value=r.get('param_enums', ''))
        ws.cell(row=row_idx, column=6, value=r.get('param_allow_null', '否'))
        ws.cell(row=row_idx, column=7, value=r.get('status', '启用'))
        ws.cell(row=row_idx, column=8, value=r.get('description', ''))
        for c in range(1, 9):
            ws.cell(row=row_idx, column=c).border = BORDER
    _style_sheet(ws)
    return wb


def build_template_template_workbook():
    """指标模板导入模板：表头 + 1 行示例 + 填写说明表。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '指标模板'
    _write_header(ws)
    sample = ['性别', '原子', '性别', 'IS_EMPTY,IS_NOT_EMPTY,IN,NOT_IN', '男,女', '否', '启用', '候选人性别']
    for c, val in enumerate(sample, start=1):
        cell = ws.cell(row=2, column=c, value=val)
        cell.fill = SAMPLE_FILL
        cell.border = BORDER
    _style_sheet(ws)

    note_ws = wb.create_sheet('填写说明')
    notes = [
        '【指标模板导入模板】',
        '1. 每行一条模板，列顺序：模板名称 / 指标类型 / 引用指标名称 / 支持的运算符 / 参数枚举 / 允许为空 / 状态 / 说明。',
        '2. 「模板名称」长度 ≤ 64，同一文件内不可重复。',
        '3. 「指标类型」填写 原子 / 派生（留空则先查原子指标再查派生指标，同名优先原子）。',
        '4. 「引用指标名称」必须已在系统中存在（原子指标或派生指标），否则该行报错。',
        '5. 「支持的运算符」填写英文 code（如 IS_EMPTY,GT）或中文（如 为空,大于），多个用逗号分隔；非法 code 报错。',
        '6. 「参数枚举」枚举型指标的可选值，多个用逗号分隔，留空表示非枚举。',
        '7. 「允许为空」填写 是 / 否（留空默认 否）。',
        '8. 「状态」填写 启用 / 停用（留空默认 启用）。',
        '9. 重复项处理由导入时的 mode 决定：skip=跳过已存在 / update=更新已存在 / error=遇重复即报错。',
        '10. 请勿修改首行表头，不要保留空白示例行以外的无关内容。',
    ]
    for i, line in enumerate(notes, start=1):
        cell = note_ws.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if i == 1:
            cell.font = Font(bold=True, size=12)
    note_ws.column_dimensions['A'].width = 100
    return wb


def build_template_error_report_workbook(original_rows, errors_by_line):
    """导入失败错误报告 xlsx。original_rows: list[{line, name, metric_kind, metric_name, operators_raw}]。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '错误明细'
    headers = ['行号', '模板名称', '指标类型', '引用指标名称', '支持的运算符', '错误原因']
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    for i, r in enumerate(original_rows, start=2):
        ws.cell(row=i, column=1, value=r.get('line'))
        ws.cell(row=i, column=2, value=r.get('name', ''))
        ws.cell(row=i, column=3, value=r.get('metric_kind', ''))
        ws.cell(row=i, column=4, value=r.get('metric_name', ''))
        ws.cell(row=i, column=5, value=r.get('operators_raw', ''))
        ws.cell(row=i, column=6, value=errors_by_line.get(r.get('line'), ''))
        for c in range(1, 7):
            ws.cell(row=i, column=c).border = BORDER
    for c, w in enumerate([8, 16, 12, 20, 28, 60], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return wb


# ===========================================================================
#  csv 构建
# ===========================================================================
def build_template_export_csv(rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(TEMPLATE_HEADERS)
    for r in rows:
        writer.writerow([
            r.get('name', ''), r.get('metric_kind', ''), r.get('metric_name', ''),
            r.get('operators', ''), r.get('param_enums', ''), r.get('param_allow_null', '否'),
            r.get('status', '启用'), r.get('description', ''),
        ])
    return buf.getvalue()


def build_template_template_csv():
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(TEMPLATE_HEADERS)
    writer.writerow(['性别', '原子', '性别', 'IS_EMPTY,IS_NOT_EMPTY,IN,NOT_IN', '男,女', '否', '启用', '候选人性别'])
    return buf.getvalue()


def build_template_error_report_csv(original_rows, errors_by_line):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(['行号', '模板名称', '指标类型', '引用指标名称', '支持的运算符', '错误原因'])
    for r in original_rows:
        writer.writerow([
            r.get('line'), r.get('name', ''), r.get('metric_kind', ''),
            r.get('metric_name', ''), r.get('operators_raw', ''),
            errors_by_line.get(r.get('line'), ''),
        ])
    return buf.getvalue()


# ===========================================================================
#  解析（xlsx / csv 自动识别）
# ===========================================================================
def parse_template_file(file_obj):
    """解析上传文件，返回 (rows, parse_errors, original_rows, errors_by_line)。

    rows: list[{name, metric_kind, metric_name, operators_raw, param_enums_raw, allow_null_raw, status_raw, description, line}]
    original_rows: list[{line, name, metric_kind, metric_name, operators_raw}]  (供错误报告)
    errors_by_line: dict[line, error_msg]  (解析级错误，如缺列)
    parse_errors: list[str]
    """
    name = (getattr(file_obj, 'name', '') or '').lower()
    if name.endswith('.csv'):
        return _parse_csv(file_obj)
    if name.endswith(('.xlsx', '.xlsm')):
        return _parse_xlsx(file_obj)
    return [], [f'不支持的文件类型：{file_obj.name}，仅支持 .xlsx / .csv'], [], {}


def _normalize_row(name, kind, metric_name, operators_raw, param_enums_raw,
                   allow_null_raw, status_raw, description, line):
    return {
        'name': (name or '').strip(),
        'metric_kind': (kind or '').strip(),
        'metric_name': (metric_name or '').strip(),
        'operators_raw': (operators_raw or '').strip(),
        'param_enums_raw': (param_enums_raw or '').strip(),
        'allow_null_raw': (allow_null_raw or '').strip(),
        'status_raw': (status_raw or '').strip(),
        'description': (description or '').strip(),
        'line': line,
    }


def _parse_xlsx(file_obj):
    rows, parse_errors, original_rows, errors_by_line = [], [], [], {}
    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — openpyxl Excel 解析异常类型不固定, 统一兜底返错误列表
        return [], [f'Excel 文件解析失败：{e}'], [], {}
    ws = wb['指标模板'] if '指标模板' in wb.sheetnames else wb.active
    iter_rows = ws.iter_rows(values_only=True)
    try:
        header = next(iter_rows)
    except StopIteration:
        return [], ['文件为空，未找到表头行'], [], {}
    header_map = {}
    for idx, h in enumerate(header):
        if h is None:
            continue
        key = str(h).strip()
        if key in TEMPLATE_HEADERS:
            header_map[key] = idx
    if '模板名称' not in header_map or '引用指标名称' not in header_map:
        return [], ['表头缺少必需的「模板名称」「引用指标名称」列'], [], {}
    col_name = header_map['模板名称']
    col_kind = header_map.get('指标类型')
    col_metric = header_map['引用指标名称']
    col_ops = header_map.get('支持的运算符')
    col_enums = header_map.get('参数枚举')
    col_null = header_map.get('允许为空')
    col_status = header_map.get('状态')
    col_desc = header_map.get('说明')

    line_no = 1  # 数据行号（与 Excel 行号对应，header 为第 1 行，数据从第 2 行起）
    for raw in iter_rows:
        line_no += 1
        name = raw[col_name] if col_name < len(raw) else None
        kind = raw[col_kind] if (col_kind is not None and col_kind < len(raw)) else ''
        metric_name = raw[col_metric] if col_metric < len(raw) else None
        operators_raw = raw[col_ops] if (col_ops is not None and col_ops < len(raw)) else ''
        param_enums_raw = raw[col_enums] if (col_enums is not None and col_enums < len(raw)) else ''
        allow_null_raw = raw[col_null] if (col_null is not None and col_null < len(raw)) else ''
        status_raw = raw[col_status] if (col_status is not None and col_status < len(raw)) else ''
        description = raw[col_desc] if (col_desc is not None and col_desc < len(raw)) else ''
        # 跳过完全空行
        if (name is None or str(name).strip() == '') and (metric_name is None or str(metric_name).strip() == ''):
            continue
        rec = _normalize_row(name, kind, metric_name, operators_raw, param_enums_raw,
                             allow_null_raw, status_raw, description, line_no)
        original_rows.append({
            'line': line_no,
            'name': rec['name'],
            'metric_kind': rec['metric_kind'],
            'metric_name': rec['metric_name'],
            'operators_raw': rec['operators_raw'],
        })
        rows.append(rec)
    return rows, parse_errors, original_rows, errors_by_line


def _parse_csv(file_obj):
    rows, parse_errors, original_rows, errors_by_line = [], [], [], {}
    try:
        content = file_obj.read()
        if isinstance(content, bytes):
            text = content.decode('utf-8-sig')
        else:
            text = content
        reader = csv.DictReader(io.StringIO(text))
        field_map = {}
        for i, f in enumerate(reader.fieldnames or []):
            key = (f or '').strip()
            if key in TEMPLATE_HEADERS:
                field_map[key] = f
        if '模板名称' not in field_map or '引用指标名称' not in field_map:
            return [], ['表头缺少必需的「模板名称」「引用指标名称」列'], [], {}
        col_name = field_map['模板名称']
        col_kind = field_map.get('指标类型')
        col_metric = field_map['引用指标名称']
        col_ops = field_map.get('支持的运算符')
        col_enums = field_map.get('参数枚举')
        col_null = field_map.get('允许为空')
        col_status = field_map.get('状态')
        col_desc = field_map.get('说明')
    except (csv.Error, ValueError, OSError) as e:  # csv 解析异常统一兜底返错误列表
        return [], [f'CSV 文件解析失败：{e}'], [], {}

    line_no = 1  # DictReader 从第 2 行起为数据（第 1 行为表头）
    for raw in reader:
        line_no += 1
        name = raw.get(col_name)
        kind = raw.get(col_kind) if col_kind else ''
        metric_name = raw.get(col_metric)
        operators_raw = raw.get(col_ops) if col_ops else ''
        param_enums_raw = raw.get(col_enums) if col_enums else ''
        allow_null_raw = raw.get(col_null) if col_null else ''
        status_raw = raw.get(col_status) if col_status else ''
        description = raw.get(col_desc) if col_desc else ''
        if (name is None or str(name).strip() == '') and (metric_name is None or str(metric_name).strip() == ''):
            continue
        rec = _normalize_row(name, kind, metric_name, operators_raw, param_enums_raw,
                             allow_null_raw, status_raw, description, line_no)
        original_rows.append({
            'line': line_no,
            'name': rec['name'],
            'metric_kind': rec['metric_kind'],
            'metric_name': rec['metric_name'],
            'operators_raw': rec['operators_raw'],
        })
        rows.append(rec)
    return rows, parse_errors, original_rows, errors_by_line


# ===========================================================================
#  审计（照搬 campus_control._log_indicator_audit 写法，entity=MetricTemplate）
# ===========================================================================
def _log_template_audit(user, action, detail, request=None, entity_id=None):
    """指标模板写操作审计（导入/导出/模板下载）。审计失败不阻断主流程。"""
    try:
        ip = (request.META.get('REMOTE_ADDR') or '') if request else ''
        ua = ((request.META.get('HTTP_USER_AGENT') or '')[:500]) if request else ''
        AuditLog.objects.create(
            user=user,
            action=action,
            entity='MetricTemplate',
            entity_id=entity_id,
            new_value=detail[:1000],
            ip=ip,
            user_agent=ua,
        )
    except (DatabaseError, ValueError, TypeError):  # 审计失败不应阻断主流程
        logger.exception('审计写入失败 entity=MetricTemplate action=%s entity_id=%s', action, entity_id)


def _template_error_payload(parse_errors, errors_by_line, original_rows, filename=''):
    """构造失败响应 data：融合错误文本 + xlsx 错误报告（base64）。"""
    errs = list(parse_errors) + [f'第 {ln} 行：{msg}' for ln, msg in sorted(errors_by_line.items())]
    payload = {
        'created': 0, 'updated': 0, 'skipped': 0, 'failed': len(errs),
        'errors': errs, 'error_file': None,
    }
    if original_rows:
        try:
            from io import BytesIO
            buf = BytesIO()
            wb = build_template_error_report_workbook(original_rows, errors_by_line)
            wb.save(buf)
            payload['error_file'] = base64.b64encode(buf.getvalue()).decode('ascii')
        except Exception as e:  # noqa: BLE001 - 报告生成失败不影响主错误返回
            logger.warning('指标模板导入错误报告生成失败: %s', e)
            payload['error_file'] = None
    return payload


class _Result:
    """服务层统一返回：(payload, status_code) 配对。View 端：Response(result.payload, status=result.status_code)。"""

    def __init__(self, payload, status_code=200):
        self.payload = payload
        self.status_code = status_code


# ===========================================================================
#  导入（照搬 campus_control.import_indicators，适配 MetricTemplate）
# ===========================================================================
def import_templates(user, file_obj, mode, filename='', request=None):
    """批量导入指标模板：数据校验 + 重复项处理（mode=skip|update|error）。

    校验：模板名称非空且 ≤64；引用指标必须存在；指标类型非法（非 原子/派生/空）→ 报错；
    运算符列按中文→code 映射归一化，非法 code → 报错；允许为空/状态按 是/否、启用/停用 解析
    （留空默认 启用 / 不允许空）；文件内同模板名不可重复。
    与库内重复按 mode 处理：skip=跳过 / update=更新 operators/param_enums/param_allow_null/status/
    description 与对应 FK / error 整批拒绝（原子回滚）。
    任一硬校验错误 → 整体 400 并附错误报告（xlsx，base64）。
    """
    if not file_obj:
        return _Result({'success': False, 'data': {
            'errors': ['缺少 file 文件字段'], 'created': 0, 'updated': 0, 'skipped': 0, 'failed': 1, 'error_file': None}}, 400)
    mode = (mode or 'skip').strip().lower()
    if mode not in ('skip', 'update', 'error'):
        return _Result({'success': False, 'data': {
            'errors': ['mode 仅支持 skip / update / error'], 'created': 0, 'updated': 0, 'skipped': 0, 'failed': 1, 'error_file': None}}, 400)

    rows, parse_errors, original_rows, errors_by_line = parse_template_file(file_obj)

    if parse_errors:
        return _Result({'success': False, 'data': _template_error_payload(parse_errors, {}, [], filename)}, 400)
    if not rows:
        return _Result({'success': False, 'data': _template_error_payload(['文件中未解析到任何有效模板行'], {}, [], filename)}, 400)

    # 预加载全部指标名（原子 / 派生），避免逐行查库
    atomic_by_name = {m.name: m for m in AtomicMetric.objects.all()}
    derived_by_name = {m.name: m for m in DerivedMetric.objects.all()}
    # 预加载已存在模板（按 name），供 mode=error 在校验阶段整批拒绝重复项
    existing_by_name = {t.name: t for t in MetricTemplate.objects.all()}

    seen_in_file = {}
    resolved = {}  # line -> {fk_field, fk_value, operators, allow_null, status, param_enums, description}
    for rec in rows:
        line = rec['line']
        name = rec['name']
        if not name:
            errors_by_line[line] = '模板名称为空'
            continue
        if len(name) > 64:
            errors_by_line[line] = f'模板名称过长（须 ≤64，当前 {len(name)}）'
            continue
        kind = rec['metric_kind']
        if kind not in ('', '原子', '派生'):
            errors_by_line[line] = f'指标类型非法：{kind!r}（填写 原子 / 派生 或留空）'
            continue
        metric_name = rec['metric_name']
        if not metric_name:
            errors_by_line[line] = '引用指标名称为空'
            continue

        # mode=error：与库内已有模板重名 → 整批拒绝（400 + 错误报告，不写库）
        if mode == 'error' and name in existing_by_name:
            errors_by_line[line] = f'模板「{name}」已存在，mode=error 不允许覆盖'
            continue

        # 解析引用指标（按指标类型列；留空则先查原子再查派生，同名优先原子）
        if kind == '原子':
            metric = atomic_by_name.get(metric_name)
            if metric is None:
                errors_by_line[line] = f'原子指标「{metric_name}」不存在'
                continue
            fk_field, fk_value = 'atomic_metric', metric
        elif kind == '派生':
            metric = derived_by_name.get(metric_name)
            if metric is None:
                errors_by_line[line] = f'派生指标「{metric_name}」不存在'
                continue
            fk_field, fk_value = 'derived_metric', metric
        else:
            atomic = atomic_by_name.get(metric_name)
            derived = derived_by_name.get(metric_name)
            if atomic is not None:
                fk_field, fk_value = 'atomic_metric', atomic
            elif derived is not None:
                fk_field, fk_value = 'derived_metric', derived
            else:
                errors_by_line[line] = f'引用指标「{metric_name}」不存在（原子 / 派生均未找到）'
                continue

        # 运算符：中文→code 归一化；非法 code → 报错；为空 → 报错
        operators, op_err = _normalize_operators(rec['operators_raw'])
        if op_err is not None:
            errors_by_line[line] = op_err
            continue
        if not operators:
            errors_by_line[line] = '未提供任何合法运算符'
            continue

        # 允许为空
        allow_null = _parse_bool(rec['allow_null_raw'])
        if allow_null is None:
            errors_by_line[line] = f'允许为空非法：{rec["allow_null_raw"]!r}（填写 是/否）'
            continue

        # 状态
        status = _parse_status(rec['status_raw'])
        if status is None:
            errors_by_line[line] = f'状态非法：{rec["status_raw"]!r}（填写 启用/停用）'
            continue

        # 参数枚举（逗号分隔 → list；留空为 []）
        param_enums = [p.strip() for p in rec['param_enums_raw'].split(',') if p.strip()]

        # 文件内重名
        if name in seen_in_file:
            errors_by_line[line] = f'与第 {seen_in_file[name]} 行重复（同名模板）'
            continue

        seen_in_file[name] = line
        resolved[line] = {
            'fk_field': fk_field,
            'fk_value': fk_value,
            'operators': operators,
            'allow_null': allow_null,
            'status': status,
            'param_enums': param_enums,
            'description': rec['description'],
        }

    if errors_by_line:
        return _Result({'success': False, 'data': _template_error_payload([], errors_by_line, original_rows, filename)}, 400)

    # 第二遍：写库（此阶段已无硬错误）
    from django.db import transaction

    created = updated = skipped = 0
    with transaction.atomic():
        for rec in rows:
            line = rec['line']
            info = resolved.get(line)
            if info is None:
                # 该校验错误行理论上已被 errors_by_line 拦截，保险跳过
                continue
            name = rec['name']
            existing = MetricTemplate.objects.filter(name=name).first()
            if existing:
                if mode == 'skip':
                    skipped += 1
                    continue
                setattr(existing, info['fk_field'], info['fk_value'])
                existing.operators = info['operators']
                existing.param_enums = info['param_enums']
                existing.param_allow_null = info['allow_null']
                existing.status = info['status']
                existing.description = info['description']
                existing.updated_by = user
                existing.save(update_fields=[
                    info['fk_field'], 'operators', 'param_enums',
                    'param_allow_null', 'status', 'description', 'updated_at',
                ])
                updated += 1
                continue
            MetricTemplate.objects.create(
                name=name,
                **{info['fk_field']: info['fk_value']},
                operators=info['operators'],
                param_enums=info['param_enums'],
                param_allow_null=info['allow_null'],
                status=info['status'],
                description=info['description'],
                created_by=user,
                updated_by=user,
            )
            created += 1

    action = 'CREATE' if created else ('UPDATE' if updated else 'READ')
    detail = f'导入指标模板完成：新建 {created} / 更新 {updated} / 跳过 {skipped}（mode={mode}）'
    _log_template_audit(user, action, detail, request=request)
    return _Result({
        'success': True,
        'data': {'created': created, 'updated': updated, 'skipped': skipped, 'failed': 0, 'errors': []},
    }, 200)
