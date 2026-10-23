"""campus_rotation — 轮岗资源导入/导出 Excel 工具（openpyxl，零额外重依赖）。

2026-10-22 兵哥指令：直播间 + 客服部门整合为「轮岗资源」一张表后，
本模块由 `io_live_room.py` 演进而来 —— 列结构从「直播间 11 列」扩为
「轮岗资源 12 列」，核心变化是**新增「资源类型」列**（直播间 / 客服部门 / 其他），
并按类型区分「平台」是否必填（2026-10-23 闸口1 裁决：**仅直播间必填**，其余类型选填）。

模板 / 导入 / 导出**共用同一套列结构**（HEADERS，顺序固定），保证
「下载模板 → 填数据 → 导入」闭环一致。

列结构（行 = 一个轮岗资源，12 列）：
  0 资源编号   resource_no  留空 = 服务端按资源类型自动补 LR+4 / SD+4 / OT+4
  1 资源名称   name         **必填**
  2 资源类型   resource_type **必填**（直播间 / 客服部门 / 其他）
  3 归属部门   department   **必填**（系统内已存在的部门名称）
  4 归属BU     bu           选填（须为二级部门 / BU）
  5 导师       mentor       选填（须具备 CAMPUS_MENTOR 身份）
  6 平台       platform     **直播间必填**（LIVE_ROOM_PLATFORM 字典，中文名或 key）；
                           客服部门 / 其他**选填**（填了同样须是启用中的字典项）
  7 地点       location     选填（**自由文本**，非字典项）
  8 容量上限   capacity     选填（非负整数）
  9 社招人数   social_count 选填（非负整数）
  10 校招人数  campus_count **只读聚合**（系统统计，导入时忽略该列）
  11 状态      is_active    启用 / 停用（留空 = 启用）

导入语义：
- 按「资源编号」upsert —— 编号留空 = 新建（服务端按类型自动补号）；
  编号已有 = 更新（命中软删行 → 跳过并提示先恢复；命中 scope 外行 → 跳过并提示无权限）。
- **单行独立校验、单行独立落库**：某行失败只跳过该行并记录原因，不中断整批。
- 导出/导入均受调用方的行级隔离（scope）约束，导出不做逐行 count（N+1 纪律）。

⚠️ 资源类型是**创建时即确定**的语义：编号前缀（LR / SD / OT）由类型决定，
且类型决定「平台」是否必填 → upsert 时若本行类型与库中既有资源类型不一致，
按**行级错误**跳过（不改类型），避免「改类型 → 编号前缀与存量号打架」。
"""
import base64
from io import BytesIO

from django.db.models import Count
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.common.excel import (
    BORDER,
    SAMPLE_FILL,
    set_column_widths,
    style_header_row,
    write_sheet,
)
from apps.core.models import Department, User
from apps.dictionary.models import DictionaryItem

from .models import RotationArrangement, RotationResource
from .validators import dict_key_active, is_bu_depth, user_has_campus_role

# ---------------------------------------------------------------------------
# 列定义（模板 / 导入 / 导出共用，顺序固定）
# ---------------------------------------------------------------------------
HEADERS = [
    '资源编号', '资源名称', '资源类型', '归属部门', '归属BU', '导师',
    '平台', '地点', '容量上限', '社招人数', '校招人数', '状态',
]
N_COLS = len(HEADERS)

# 导出行的取值顺序（与 HEADERS 一一对应）
_ROW_KEYS = [
    'resource_no', 'name', 'resource_type', 'department', 'bu', 'mentor',
    'platform', 'location', 'capacity', 'social_count', 'campus_count', 'status',
]

# 表头定位用的锚点列（含它的那一行即表头行）
_HEADER_ANCHOR = '资源名称'
# 必填列（缺任一列 → 文件级错误，无法导入）
_REQUIRED_HEADERS = ('资源名称', '资源类型', '归属部门')
# 扫描表头行的最大行数（允许说明区在上方）
_HEADER_SCAN_ROWS = 30

# 样式（领域专有；表头/边框/示例底纹复用 apps.common.excel）
NOTE_FILL = PatternFill('solid', fgColor='FFF7ED')
NOTE_FONT = Font(color='9A3412', size=10)
ERROR_FILL = PatternFill('solid', fgColor='FEE2E2')  # 浅红底标注错误行
ERROR_FONT = Font(color='B91C1C', size=11)

ERROR_COL = '错误原因'

COLUMN_WIDTHS = [14, 22, 12, 16, 16, 14, 12, 14, 10, 10, 10, 8]

STATUS_ACTIVE = '启用'
STATUS_INACTIVE = '停用'

PLATFORM_DICT = 'LIVE_ROOM_PLATFORM'

# 资源类型：中文名 <-> 枚举 key 双向映射（导入接受两种写法）
RESOURCE_TYPE_BY_LABEL = {
    '直播间': 'LIVE_ROOM',
    '客服部门': 'SERVICE_DEPT',
    '其他': 'OTHER',
}
RESOURCE_TYPE_LABEL = {
    'LIVE_ROOM': '直播间',
    'SERVICE_DEPT': '客服部门',
    'OTHER': '其他',
}

# 字段长度上限（与模型 max_length 对齐）—— 超出走行级错误，不撞数据库约束
_MAX_NAME_LEN = 128        # RotationResource.name
_MAX_NO_LEN = 32           # RotationResource.resource_no
_MAX_LOCATION_LEN = 64     # RotationResource.location
_MAX_PLATFORM_LEN = 32     # RotationResource.platform


# ---------------------------------------------------------------------------
# 小工具
# ---------------------------------------------------------------------------
def _s(v):
    """单元格值 → 去空格的字符串（None → ''；整数型浮点不落 .0）。"""
    if v is None:
        return ''
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _parse_nonneg_int(raw, label):
    """解析非负整数 → (值, 错误信息)。空 → (None, None)；非法/负数 → (None, 错误信息)。"""
    text = _s(raw)
    if not text:
        return None, None
    try:
        num = int(round(float(text)))
    except (TypeError, ValueError):
        return None, f'{label}「{text}」须为非负整数'
    if num < 0:
        return None, f'{label}「{text}」不能为负'
    return num, None


def _user_display(user):
    """用户显示名：优先姓名，其次 username，都没有则空串。"""
    if user is None:
        return ''
    return user.get_full_name() or user.username or ''


def _dict_display_names(type_code):
    """{字典 key: 字典中文名}（含停用项 —— 导出只做展示，停用项也照常显示名字）。"""
    rows = DictionaryItem.objects.filter(type__code=type_code).values('key', 'value')
    return {r['key']: r['value'] for r in rows}


def _dict_items(type_code):
    """{(key, value)} → 字典项 dict（含停用项；停用项由 dict_key_active 拒绝）。"""
    by_key, by_value = {}, {}
    rows = DictionaryItem.objects.filter(type__code=type_code).values('key', 'value')
    for r in rows:
        by_key.setdefault(r['key'], r)
        by_value.setdefault(r['value'], r)
    return by_key, by_value


def _resolve_dict_key(type_code, raw, cache, active_cache):
    """把「中文名或 key」解析为启用字典项的 key；解析不到（含停用项）返回 None。

    复用 `validators.dict_key_active` 做「存在且启用」的判定（按 (type, key) 缓存，
    避免逐行重复查库）；中文名 → key 的映射一次批量取回，避免 N+1。
    """
    by_key, by_value = cache[type_code]
    item = by_key.get(raw) or by_value.get(raw)
    if item is None:
        return None
    memo_key = (type_code, item['key'])
    if memo_key not in active_cache:
        active_cache[memo_key] = dict_key_active(type_code, item['key'])
    if not active_cache[memo_key]:
        return None
    return item['key']


def _parse_resource_type(raw):
    """资源类型单元格 → (枚举 key, 错误信息)。接受中文名或枚举 key。"""
    text = _s(raw)
    if not text:
        return None, '资源类型为空'
    if text in RESOURCE_TYPE_LABEL:
        return text, None
    if text in RESOURCE_TYPE_BY_LABEL:
        return RESOURCE_TYPE_BY_LABEL[text], None
    return None, (
        f'资源类型「{text}」非法（须为 直播间 / 客服部门 / 其他）'
    )


# ---------------------------------------------------------------------------
# 导出
# ---------------------------------------------------------------------------
def _campus_counts():
    """{资源 id: 在校(ONGOING)占用人数} —— 一次聚合查询（N+1 纪律）。

    口径与 `RotationResourceSerializer.to_representation` 一致：按**资源**
    （`RotationArrangement.resource_id`）+ `stage='ONGOING'` 统计。
    """
    rows = (
        RotationArrangement.objects.filter(stage='ONGOING', deleted_at__isnull=True)
        .exclude(resource_id__isnull=True)
        .values('resource_id')
        .annotate(c=Count('id'))
    )
    return {r['resource_id']: r['c'] for r in rows}


def build_export_rows(qs):
    """轮岗资源 queryset → 导出用的展示行（列表）。

    · 资源类型写**中文名**（直播间 / 客服部门 / 其他）；
    · 平台写**中文名**（字典 value）；地点为**自由文本**原样输出；
    · 状态写 启用/停用；导师写显示名；
    · 归属部门 / 归属BU 写 Department.name；
    · 校招人数为一次性聚合结果（不做逐行 count）。
    """
    resources = list(qs.select_related('department', 'bu', 'mentor').order_by('resource_no'))
    counts = _campus_counts()
    platform_names = _dict_display_names(PLATFORM_DICT)

    rows = []
    for r in resources:
        rows.append({
            'resource_no': r.resource_no or '',
            'name': r.name or '',
            'resource_type': RESOURCE_TYPE_LABEL.get(r.resource_type, r.resource_type or ''),
            'department': r.department.name if r.department_id else '',
            'bu': r.bu.name if r.bu_id else '',
            'mentor': _user_display(r.mentor) if r.mentor_id else '',
            'platform': platform_names.get(r.platform, r.platform or ''),
            'location': r.location or '',
            'capacity': r.capacity,
            'social_count': r.social_count,
            'campus_count': counts.get(r.id, 0),
            'status': STATUS_ACTIVE if r.is_active else STATUS_INACTIVE,
        })
    return rows


def build_export_workbook(rows):
    """rows: list[dict]（build_export_rows 的产物）→ Workbook。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '轮岗资源'
    data = [[r.get(k, '') for k in _ROW_KEYS] for r in rows]
    write_sheet(ws, HEADERS, data, column_widths=COLUMN_WIDTHS)
    return wb


# ---------------------------------------------------------------------------
# 模板
# ---------------------------------------------------------------------------
TEMPLATE_NOTES = [
    '填写说明：',
    '1. 每行一个轮岗资源；列顺序：资源编号 / 资源名称 / 资源类型 / 归属部门 / 归属BU / '
    '导师 / 平台 / 地点 / 容量上限 / 社招人数 / 校招人数 / 状态。',
    '2. 必填列：资源名称、资源类型、归属部门；其余列留空即不填。',
    '3. 资源类型填「直播间」「客服部门」或「其他」；资源编号留空 = 由系统按类型自动生成'
    '（直播间 LR+4 位、客服部门 SD+4 位、其他 OT+4 位，如 LR0001）；'
    '填写则按该编号 upsert（命中已有编号 → 更新该资源）。',
    '4. 归属部门必须填写系统内已存在的部门名称；归属BU 填写时必须是二级部门（BU）。',
    '5. 平台：直播间必填，客服部门 / 其他选填；可填中文名（如 视频号）或字典 key（如 '
    'VIDEO_ACCOUNT），两者都认；填了就须为启用状态的平台字典项，留空表示不填。',
    '6. 地点为自由文本，直接填写即可（如 北京 / 滨海新区 A 座），不做字典校验。',
    '7. 导师可填用户名或姓名，该用户须具备「校园导师（CAMPUS_MENTOR）」身份。',
    '8. 容量上限 / 社招人数填非负整数；状态填「启用」或「停用」（留空 = 启用）。',
    '9. 校招人数由系统按在轮岗（ONGOING）人数统计，导入时忽略该列。',
    '10. 单行校验失败只跳过该行并记录原因，不影响其余行；错误明细见导入返回的「错误明细」文件。',
    '11. 请勿修改表头行与下方示例行以外的结构；导入按表头名定位列，列顺序可自行调整。',
]

TEMPLATE_SAMPLE = [
    'LR0001', '示例直播间', '直播间', '技术BU', '技术BU', '张三',
    '视频号', '北京', 30, 5, '', STATUS_ACTIVE,
]


def build_template_workbook():
    """返回（带说明区 + 表头 + 1 行示例的模板 Workbook, 表头行号）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '轮岗资源模板'

    # 1) 说明区
    for i, line in enumerate(TEMPLATE_NOTES, start=1):
        cell = ws.cell(row=i, column=1, value=line)
        cell.fill = NOTE_FILL
        cell.font = NOTE_FONT
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=N_COLS)
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # 2) 表头（说明区后空一行）
    header_row = len(TEMPLATE_NOTES) + 2
    write_sheet(ws, HEADERS, [], start_row=header_row, column_widths=COLUMN_WIDTHS)

    # 3) 示例行
    for c, v in enumerate(TEMPLATE_SAMPLE, start=1):
        cell = ws.cell(row=header_row + 1, column=c, value=v)
        cell.fill = SAMPLE_FILL
        cell.border = BORDER

    set_column_widths(ws, COLUMN_WIDTHS)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    return wb, header_row


# ---------------------------------------------------------------------------
# 导入解析
# ---------------------------------------------------------------------------
def _dept_maps():
    """(按名称, 按编码) → Department；同名取先创建者。"""
    by_name, by_code = {}, {}
    for d in Department.objects.all():
        by_name.setdefault(d.name, d)
        by_code.setdefault(d.code, d)
    return by_name, by_code


def _mentor_map():
    """{用户名 或 姓名} → 用户 dict（id / username / first_name / last_name）。"""
    out = {}
    rows = User.objects.filter(is_active=True).values(
        'id', 'username', 'first_name', 'last_name',
    )
    for u in rows:
        out.setdefault(u['username'], u)
        full = f"{u['first_name'] or ''} {u['last_name'] or ''}".strip()
        if full:
            out.setdefault(full, u)
    return out


def parse_import_workbook(file_obj):
    """解析上传的 xlsx → (rows, errors, original_rows, errors_by_line)。

    rows: list[合法行 payload]（已通过单行校验，可直接 upsert）。
    errors: list[str]（行级错误文案，已带「第N行：」前缀，给用户看）。
    original_rows: list[(line_no, 原始单元格值列表)]（HEADERS 顺序，供错误报告用）。
    errors_by_line: dict[line_no → 错误原因]。
    """
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active

    max_col = ws.max_column or N_COLS
    header_row = None
    for ridx in range(1, min(ws.max_row, _HEADER_SCAN_ROWS) + 1):
        vals = [_s(ws.cell(row=ridx, column=c).value) for c in range(1, max_col + 1)]
        if _HEADER_ANCHOR in vals:
            header_row = ridx
            break
    if header_row is None:
        return [], [f'未找到表头（需含「{_HEADER_ANCHOR}」列）'], [], {}

    header_map = {}
    for c in range(1, max_col + 1):
        name = _s(ws.cell(row=header_row, column=c).value)
        if name and name not in header_map:
            header_map[name] = c - 1
    missing = [h for h in _REQUIRED_HEADERS if h not in header_map]
    if missing:
        return [], [f'表头缺少必需的「{"」「".join(missing)}」列'], [], {}

    # 一次性取回校验所需的字典 / 部门 / 用户（避免逐行查库）
    dept_by_name, dept_by_code = _dept_maps()
    mentor_map = _mentor_map()
    dict_cache = {PLATFORM_DICT: _dict_items(PLATFORM_DICT)}
    dict_active_cache = {}

    rows = []
    errors = []
    original_rows = []
    errors_by_line = {}

    for ridx in range(header_row + 1, ws.max_row + 1):
        raw_cells = [ws.cell(row=ridx, column=c).value for c in range(1, max_col + 1)]
        vals = {
            name: (raw_cells[idx] if idx < len(raw_cells) else None)
            for name, idx in header_map.items()
        }
        if all(not _s(vals.get(h)) for h in HEADERS):
            continue  # 空行
        line_no = ridx
        original_rows.append((line_no, [vals.get(h) for h in HEADERS]))

        def _row_err(msg):
            errors.append(f'第{line_no}行：{msg}')
            errors_by_line[line_no] = msg

        # —— 资源名称（必填）——
        name = _s(vals.get('资源名称'))
        if not name:
            _row_err('资源名称为空')
            continue
        if len(name) > _MAX_NAME_LEN:
            _row_err(f'资源名称过长（须 ≤{_MAX_NAME_LEN} 字，当前 {len(name)}）')
            continue

        # —— 资源类型（必填；直播间 / 客服部门 / 其他）——
        resource_type, type_err = _parse_resource_type(vals.get('资源类型'))
        if type_err:
            _row_err(type_err)
            continue

        # —— 归属部门（必填，须系统内已存在）——
        dept_raw = _s(vals.get('归属部门'))
        if not dept_raw:
            _row_err('归属部门为空')
            continue
        dept = dept_by_name.get(dept_raw)
        if dept is None:
            _row_err(f'归属部门「{dept_raw}」不存在（须填写系统内已存在的部门名称）')
            continue
        if not dept.is_active:
            _row_err(f'归属部门「{dept_raw}」未启用')
            continue

        # —— 归属BU（选填；须存在 + 启用 + 二级部门）——
        bu_raw = _s(vals.get('归属BU'))
        bu = None
        if bu_raw:
            bu = dept_by_name.get(bu_raw) or dept_by_code.get(bu_raw)
            if bu is None:
                _row_err(f'归属BU「{bu_raw}」不存在')
                continue
            if not bu.is_active:
                _row_err(f'归属BU「{bu_raw}」未启用')
                continue
            if not is_bu_depth(bu):
                _row_err(f'归属BU「{bu_raw}」须为二级部门（部门层级=第2级）')
                continue

        # —— 平台（直播间必填；其余类型选填）——
        # 2026-10-23 闸口1 裁决 Q1：直播间 / 客服部门 / 其他 **均可填平台**，
        #   仅「直播间」保留必填判空；有值则一律须解析为启用中的字典项。
        platform_raw = _s(vals.get('平台'))
        platform_key = ''
        if resource_type == 'LIVE_ROOM' and not platform_raw:
            _row_err('资源类型为「直播间」时，平台不能为空')
            continue
        if platform_raw:
            platform_key = _resolve_dict_key(
                PLATFORM_DICT, platform_raw, dict_cache, dict_active_cache,
            )
            if not platform_key:
                _row_err(
                    f'平台「{platform_raw}」非法（须为启用中的平台字典项，可填中文名或 key）'
                )
                continue
        if len(platform_key) > _MAX_PLATFORM_LEN:
            _row_err(f'平台过长（须 ≤{_MAX_PLATFORM_LEN} 字，当前 {len(platform_key)}）')
            continue

        # —— 地点（选填；自由文本，原样写入，不做字典校验）——
        location_raw = _s(vals.get('地点'))
        if len(location_raw) > _MAX_LOCATION_LEN:
            _row_err(f'地点过长（须 ≤{_MAX_LOCATION_LEN} 字，当前 {len(location_raw)}）')
            continue

        # —— 容量上限 / 社招人数（非负整数）——
        capacity, err = _parse_nonneg_int(vals.get('容量上限'), '容量上限')
        if err:
            _row_err(err)
            continue
        social_count, err = _parse_nonneg_int(vals.get('社招人数'), '社招人数')
        if err:
            _row_err(err)
            continue

        # —— 导师（选填；须具备 CAMPUS_MENTOR 身份）——
        mentor_raw = _s(vals.get('导师'))
        mentor_id = None
        if mentor_raw:
            mentor = mentor_map.get(mentor_raw)
            if mentor is None:
                _row_err(f'导师「{mentor_raw}」不存在')
                continue
            if not user_has_campus_role(mentor['id'], 'CAMPUS_MENTOR'):
                _row_err(f'导师「{mentor_raw}」不具备校园导师（CAMPUS_MENTOR）身份')
                continue
            mentor_id = mentor['id']

        # —— 状态（空 = 启用 / 启用 / 停用）——
        status_raw = _s(vals.get('状态'))
        if status_raw in ('', STATUS_ACTIVE):
            is_active = True
        elif status_raw == STATUS_INACTIVE:
            is_active = False
        else:
            _row_err(f'状态「{status_raw}」非法（须为 启用 / 停用，或留空=启用）')
            continue

        # 校招人数 = 只读聚合列，导入时忽略（不校验、不写入）；备注列已下线，导入不写
        resource_no = _s(vals.get('资源编号'))
        if len(resource_no) > _MAX_NO_LEN:
            _row_err(f'资源编号过长（须 ≤{_MAX_NO_LEN} 字，当前 {len(resource_no)}）')
            continue
        rows.append({
            'line': line_no,
            'resource_no': resource_no,
            'name': name,
            'resource_type': resource_type,
            'department': dept,
            'bu': bu,
            'mentor_id': mentor_id,
            'platform': platform_key,
            'location': location_raw,
            'capacity': capacity,
            'social_count': social_count if social_count is not None else 0,
            'is_active': is_active,
        })

    return rows, errors, original_rows, errors_by_line


# ---------------------------------------------------------------------------
# 导入落库（upsert）
# ---------------------------------------------------------------------------
def import_resources(rows, user, scope_qs):
    """按「资源编号」upsert 轮岗资源（逐行独立落库，某行失败只跳过该行）。

    rows: parse_import_workbook 产出的合法行 payload。
    user: 当前登录用户（写 created_by / updated_by）。
    scope_qs: callable(qs) -> qs —— 调用方的行级隔离口径
              （RotationResourceViewSet.scope_queryset，scope_field='bu'）。
    返回 (created, updated, errors, errors_by_line)。

    ⚠️ 「未导入行数」不在此返回 —— 调用方按 `len(errors_by_line)` 统计
    （校验失败行 + 本阶段的 scope 外 / 已删除跳过行，每行恰好一条错误消息）。
    """
    created = updated = 0
    errors = []
    errors_by_line = {}

    resource_nos = [r['resource_no'] for r in rows if r['resource_no']]
    existing = {}
    scope_ids = set()
    if resource_nos:
        # 含软删行：命中已软删行要提示「已删除，请先恢复」，不能当成「查不到 → 新建」
        existing = {
            o.resource_no: o for o in RotationResource.all_objects.filter(
                resource_no__in=resource_nos,
            )
        }
        # 行级隔离：scope 外的资源不可 upsert
        scope_ids = set(
            scope_qs(RotationResource.all_objects.filter(resource_no__in=resource_nos))
            .values_list('id', flat=True)
        )

    for r in rows:
        line_no = r['line']
        target = existing.get(r['resource_no']) if r['resource_no'] else None

        if target is not None:
            if target.id not in scope_ids:
                msg = f'轮岗资源「{target.resource_no}」不在你的数据范围内，无权限操作该资源'
                errors.append(f'第{line_no}行：{msg}')
                errors_by_line[line_no] = msg
                continue
            if target.deleted_at is not None:
                msg = f'轮岗资源「{target.resource_no}」已删除，请先恢复后再导入'
                errors.append(f'第{line_no}行：{msg}')
                errors_by_line[line_no] = msg
                continue
            if target.resource_type != r['resource_type']:
                msg = (
                    f'轮岗资源「{target.resource_no}」的资源类型为'
                    f'「{RESOURCE_TYPE_LABEL.get(target.resource_type, target.resource_type)}」，'
                    f'与本行的「{RESOURCE_TYPE_LABEL[r["resource_type"]]}」不一致'
                )
                errors.append(f'第{line_no}行：{msg}')
                errors_by_line[line_no] = msg
                continue
            target.name = r['name']
            target.department = r['department']
            target.bu = r['bu']
            target.mentor_id = r['mentor_id']
            target.platform = r['platform']
            target.location = r['location']
            target.capacity = r['capacity']
            target.social_count = r['social_count']
            target.is_active = r['is_active']
            target.updated_by = user
            # 备注列已下线 → 导入不写（保留库里既有备注）
            target.save(update_fields=[
                'name', 'department', 'bu', 'mentor', 'platform', 'location',
                'capacity', 'social_count', 'is_active',
                'updated_by', 'updated_at',
            ])
            updated += 1
            continue

        # 新建：编号留空 → RotationResource.save() 按类型自动补 LR+4 / SD+4 / OT+4
        RotationResource.objects.create(
            resource_no=r['resource_no'],
            name=r['name'],
            resource_type=r['resource_type'],
            department=r['department'],
            bu=r['bu'],
            mentor_id=r['mentor_id'],
            platform=r['platform'],
            location=r['location'],
            capacity=r['capacity'],
            social_count=r['social_count'],
            is_active=r['is_active'],
            created_by=user,
            updated_by=user,
        )
        created += 1

    return created, updated, errors, errors_by_line


# ---------------------------------------------------------------------------
# 错误报告
# ---------------------------------------------------------------------------
def build_error_report_workbook(original_rows, errors_by_line):
    """生成「导入错误明细」xlsx：保留原始数据行，并在每行末尾追加「错误原因」列。

    original_rows: list[(line_no, row_values)]（来自 parse_import_workbook）。
    errors_by_line: dict[line_no → 错误原因字符串]。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = '导入错误明细'
    headers = list(HEADERS) + [ERROR_COL]
    write_sheet(ws, headers, [], column_widths=COLUMN_WIDTHS + [50])
    style_header_row(ws, row=1)
    # 错误原因列用专属红底样式（覆盖通用表头）
    err_header = ws.cell(row=1, column=len(headers))
    err_header.fill = ERROR_FILL
    err_header.font = ERROR_FONT

    for line_no, vals in original_rows:
        err = errors_by_line.get(line_no, '')
        ws.append(list(vals) + [err])
        excel_row = ws.max_row
        if err:
            for c in range(1, len(headers) + 1):
                ws.cell(row=excel_row, column=c).fill = ERROR_FILL
            ws.cell(row=excel_row, column=len(headers)).font = ERROR_FONT

    set_column_widths(ws, COLUMN_WIDTHS + [50])
    ws.freeze_panes = 'A2'
    return wb


def build_error_report_base64(original_rows, errors_by_line):
    """错误报告 xlsx → base64 字符串（无原始行时返回 None）。

    报告生成失败**不阻断**主错误返回（fail-soft）——异常集合收窄到
    openpyxl/IO 层的具体异常，不新增宽捕获（见 tests/test_no_blind_except.py）。
    """
    if not original_rows:
        return None
    try:
        buf = BytesIO()
        build_error_report_workbook(original_rows, errors_by_line).save(buf)
        return base64.b64encode(buf.getvalue()).decode('ascii')
    except (OSError, ValueError, TypeError, KeyError):
        # 报告生成失败不影响主错误返回
        return None


__all__ = [
    'HEADERS', 'N_COLS', 'ERROR_COL', 'COLUMN_WIDTHS',
    'RESOURCE_TYPE_BY_LABEL', 'RESOURCE_TYPE_LABEL',
    'build_export_rows', 'build_export_workbook', 'build_template_workbook',
    'parse_import_workbook', 'import_resources',
    'build_error_report_workbook', 'build_error_report_base64',
]
