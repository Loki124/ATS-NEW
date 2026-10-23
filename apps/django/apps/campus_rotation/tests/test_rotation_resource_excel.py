"""轮岗资源 Excel 导入/导出 + 「按类型前缀取号」规则 回归测试（2026-10-22 兵哥指令）。

覆盖（直播间 + 客服部门整合为「轮岗资源」一张表后）：
1. 编号：按**资源类型前缀独立计数**（直播间 LR+4 / 客服部门 SD+4 / 其他 OT+4）；
   软删行占号不复用；显式 resource_no 不被覆盖；
2. 导出：200 + xlsx Content-Type + 12 列表头 + 行数 = 可见资源数 +
   平台写中文名 + 资源类型写中文名；
3. 模板：200 + xlsx 可读 + 表头含「资源名称」；
4. 导入：无编号 → 新建 + 自动按类型补号；已有编号 → 更新；
   部门 / 资源类型 / 平台 / 容量 / 状态非法 → 跳过 + 「第N行」错误文案；
   校招人数列被忽略；地点为自由文本（不校验字典）；
   软删行 / scope 外行 → 跳过；类型与既有资源不一致 → 跳过；
5. 行级隔离：scope 外用户的导出不含他人资源。

fixtures（auth_client 等）来自 tests/fixtures_common（经顶层 conftest.py 以 plugin 全局加载）。
"""
import io

import pytest
from django.contrib.auth import get_user_model
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from apps.campus_rotation.io_rotation_resource import HEADERS
from apps.campus_rotation.models import RotationResource
from apps.campus_rotation.seeds import (
    seed_campus_roles,
    seed_live_room_platform_type,
    seed_location_type,
)
from apps.core.models import Department
from apps.core.models_permission_v2 import RolePermissionV2, RoleV2, UserRoleV2

BASE = '/api/v1/campus-rotation'
XLSX_CT = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
USER = get_user_model()


# --------------------------------------------------------------------------- #
# fixtures
# --------------------------------------------------------------------------- #
@pytest.fixture
def _seed_campus_data(db):
    """确保校招轮岗角色与字典种子就绪（测试库幂等）。"""
    seed_campus_roles()
    seed_live_room_platform_type()
    seed_location_type()


@pytest.fixture
def bu_hierarchy(db):
    """部门树（0 基层级）：root(0) -> grp(1, 事业群) -> bu_valid(2, BU) / bu_other(2, BU)。"""
    root = Department.objects.create(id='dept-root', name='集团', code='ROOT', path='/集团')
    grp = Department.objects.create(
        id='dept-grp', name='技术事业群', code='GRP', parent=root, path='/集团/技术事业群',
    )
    bu_valid = Department.objects.create(
        id='dept-bu', name='技术BU', code='BU', parent=grp, path='/集团/技术事业群/技术BU',
    )
    bu_other = Department.objects.create(
        id='dept-bu2', name='市场BU', code='BU2', parent=grp, path='/集团/技术事业群/市场BU',
    )
    return {'root': root, 'grp': grp, 'bu_valid': bu_valid, 'bu_other': bu_other}


@pytest.fixture
def limited_client(db, bu_hierarchy):
    """范围受限的普通用户：DEPT 范围（仅本部门）+ 持 resource:list / resource:create 码。

    其可见范围 = 归属BU 落在「技术BU」的轮岗资源；归属BU =「市场BU」的不可见、不可改。
    """
    RoleV2.objects.create(
        system_code='recruit', role_code='CAMPUS_RESOURCE_VIEWER',
        role_name='轮岗资源查看', default_data_scope_type='DEPT', is_system=1, status=1,
    )
    for code in ('recruit:campus:resource:list', 'recruit:campus:resource:create'):
        RolePermissionV2.objects.create(
            role_code='CAMPUS_RESOURCE_VIEWER', resource_code=code, system_code='recruit',
        )
    user = USER.objects.create_user(
        username='limited_viewer', password='Test@1234', employee_id='L-001',
        department=bu_hierarchy['bu_valid'],
    )
    UserRoleV2.objects.create(
        user=user, role_code='CAMPUS_RESOURCE_VIEWER', system_code='recruit',
    )
    client = APIClient()
    refresh = RefreshToken.for_user(user)
    client.credentials(HTTP_AUTHORIZATION=f'Bearer {refresh.access_token}')
    return client


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _reset_resources():
    """清空轮岗资源（含软删行），让每个用例从空表起步。"""
    RotationResource.all_objects.all().delete()


def _row(resource_no='', name='', resource_type='直播间', department='', bu='', mentor='',
         platform='', location='', capacity='', social='', campus='', status=''):
    """按 HEADERS 顺序拼一行（空串 = 该列不填）。"""
    return [
        resource_no, name, resource_type, department, bu, mentor, platform,
        location, capacity, social, campus, status,
    ]


def _xlsx_bytes(rows):
    """rows: list[list]（每行长度与 HEADERS 一致）→ xlsx 字节流。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '轮岗资源'
    ws.append(list(HEADERS))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(content, name='rotation_resources.xlsx'):
    from django.core.files.uploadedfile import SimpleUploadedFile
    return SimpleUploadedFile(name, content, content_type=XLSX_CT)


def _post_import(client, rows):
    content = _xlsx_bytes(rows)
    return client.post(
        f'{BASE}/resources/import/',
        {'file': _upload(content)},
        format='multipart',
    )


def _export_rows(client):
    resp = client.get(f'{BASE}/resources/export/?file_format=xlsx')
    assert resp.status_code == 200, resp.content
    ws = load_workbook(io.BytesIO(resp.content))['轮岗资源']
    return list(ws.iter_rows(min_row=2, values_only=True))


# --------------------------------------------------------------------------- #
# 1. 编号规则（按类型前缀独立计数）
# --------------------------------------------------------------------------- #
@pytest.mark.django_db
class TestResourceNoRule:
    def test_prefix_depends_on_resource_type(self):
        """直播间 LR+4 / 客服部门 SD+4 / 其他 OT+4，且三个序列**互相独立**。"""
        _reset_resources()
        live = RotationResource.objects.create(
            name='直播间A', resource_type='LIVE_ROOM', platform='VIDEO_ACCOUNT',
        )
        service = RotationResource.objects.create(
            name='客服部A', resource_type='SERVICE_DEPT',
        )
        other = RotationResource.objects.create(
            name='其它A', resource_type='OTHER',
        )
        assert live.resource_no == 'LR0001', live.resource_no
        assert service.resource_no == 'SD0001', service.resource_no
        assert other.resource_no == 'OT0001', other.resource_no

        # 直播间再建 → 仍走 LR 序列（不被 SD/OT 干扰）
        live2 = RotationResource.objects.create(
            name='直播间B', resource_type='LIVE_ROOM', platform='VIDEO_ACCOUNT',
        )
        assert live2.resource_no == 'LR0002', live2.resource_no

    def test_soft_deleted_row_still_occupies_number(self):
        """🔴 软删行占号：唯一索引覆盖软删行 → 取号须看全集，否则复用被删行的号必撞约束。"""
        _reset_resources()
        a = RotationResource.objects.create(
            name='将被软删', resource_type='LIVE_ROOM', platform='VIDEO_ACCOUNT',
        )
        a.soft_delete()
        assert RotationResource.all_objects.filter(resource_no='LR0001').exists(), \
            '软删行应仍在库里占号'
        b = RotationResource.objects.create(
            name='软删后新建', resource_type='LIVE_ROOM', platform='VIDEO_ACCOUNT',
        )
        assert b.resource_no != 'LR0001', '复用了软删行占用的号 → 必然撞唯一索引'
        assert b.resource_no == 'LR0002', b.resource_no

    def test_explicit_resource_no_is_not_overwritten(self):
        """显式给定 resource_no 时不被自动补号覆盖。"""
        _reset_resources()
        legacy = RotationResource.objects.create(
            resource_no='ZB000123', name='存量资源', resource_type='LIVE_ROOM',
            platform='VIDEO_ACCOUNT',
        )
        assert legacy.resource_no == 'ZB000123', legacy.resource_no
        fresh = RotationResource.objects.create(
            name='新号资源', resource_type='LIVE_ROOM', platform='VIDEO_ACCOUNT',
        )
        assert fresh.resource_no == 'LR0001', fresh.resource_no


# --------------------------------------------------------------------------- #
# 2. 导出
# --------------------------------------------------------------------------- #
@pytest.mark.django_db
class TestResourceExport:
    def test_export_returns_xlsx_with_twelve_columns(self, auth_client, bu_hierarchy, _seed_campus_data):
        _reset_resources()
        bu = bu_hierarchy['bu_valid']
        RotationResource.objects.create(
            name='直播间A', resource_type='LIVE_ROOM', department=bu, bu=bu,
            platform='VIDEO_ACCOUNT', location='BEIJING',
        )
        RotationResource.objects.create(
            name='客服部A', resource_type='SERVICE_DEPT', department=bu, bu=bu,
        )
        resp = auth_client.get(f'{BASE}/resources/export/?file_format=xlsx')
        assert resp.status_code == 200, resp.content
        assert resp['Content-Type'] == XLSX_CT
        assert 'rotation_resources_export.xlsx' in resp['Content-Disposition']

        ws = load_workbook(io.BytesIO(resp.content))['轮岗资源']
        header = [c.value for c in ws[1]]
        assert header == HEADERS, header
        assert len(header) == 12, f'应为 12 列，实际 {len(header)}'

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 2, f'行数应等于可见资源数（2），实际 {len(rows)}'
        # 资源类型写中文名
        assert {r[2] for r in rows} == {'直播间', '客服部门'}
        # 平台写中文名（字典 value），不写 key；空单元格回读为 None → 归一为空串
        assert {r[6] or '' for r in rows} == {'视频号', ''}
        # 状态写 启用 / 停用
        assert {r[11] for r in rows} == {'启用'}

    def test_export_file_format_defaults_to_xlsx(self, auth_client):
        resp = auth_client.get(f'{BASE}/resources/export/')
        assert resp.status_code == 200, resp.content
        assert resp['Content-Type'] == XLSX_CT


# --------------------------------------------------------------------------- #
# 3. 模板
# --------------------------------------------------------------------------- #
@pytest.mark.django_db
class TestResourceTemplate:
    def test_template_is_readable_xlsx_with_header(self, auth_client):
        resp = auth_client.get(f'{BASE}/resources/template/?file_format=xlsx')
        assert resp.status_code == 200, resp.content
        assert resp['Content-Type'] == XLSX_CT
        assert 'rotation_resources_template.xlsx' in resp['Content-Disposition']

        ws = load_workbook(io.BytesIO(resp.content)).active
        header_row = None
        for ridx in range(1, min(ws.max_row, 30) + 1):
            vals = [
                str(ws.cell(row=ridx, column=c).value or '').strip()
                for c in range(1, len(HEADERS) + 1)
            ]
            if '资源名称' in vals:
                header_row = ridx
                break
        assert header_row is not None, '模板应能定位到表头行（含「资源名称」）'
        header = [ws.cell(row=header_row, column=c).value for c in range(1, len(HEADERS) + 1)]
        assert header == HEADERS, header

    def test_template_notes_platform_rule_covers_all_types(self, auth_client, bu_hierarchy, _seed_campus_data):
        """2026-10-23 闸口1：模板填写说明的平台口径 = 直播间必填 + 其余类型选填，且不含旧口径文案。"""
        resp = auth_client.get(f'{BASE}/resources/template/?file_format=xlsx')
        assert resp.status_code == 200, resp.content
        ws = load_workbook(io.BytesIO(resp.content)).active
        texts = []
        for ridx in range(1, min(ws.max_row, 30) + 1):
            for cidx in range(1, len(HEADERS) + 1):
                v = ws.cell(row=ridx, column=cidx).value
                if isinstance(v, str) and '平台' in v:
                    texts.append(v)
        joined = '\n'.join(texts)
        assert joined, '模板填写说明应含平台相关文案'
        assert '选填' in joined, joined
        assert '客服部门' in joined, joined
        # 旧口径「…其余类型请留空」已作废（闸口1 Q1：三类形态均可填平台）
        assert '请留空' not in joined, joined


# --------------------------------------------------------------------------- #
# 4. 导入
# --------------------------------------------------------------------------- #
@pytest.mark.django_db
class TestResourceImport:
    def test_import_without_resource_no_creates_and_auto_numbers(self, auth_client, bu_hierarchy, _seed_campus_data):
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(
                name='导入直播间', resource_type='直播间', department='技术BU', bu='技术BU',
                platform='视频号', location='北京', capacity=20, social=3, status='启用',
            ),
        ])
        assert resp.status_code == 200, resp.content
        body = resp.json()
        assert body['success'] is True
        assert body['data']['created'] == 1
        assert body['data']['updated'] == 0
        assert body['data']['skipped'] == 0

        res = RotationResource.objects.get(name='导入直播间')
        assert res.resource_no == 'LR0001', res.resource_no
        assert res.resource_type == 'LIVE_ROOM'
        assert res.platform == 'VIDEO_ACCOUNT', '平台中文名应解析为字典 key'
        assert res.location == '北京', '地点为自由文本 → 原样入库'
        assert res.department_id == 'dept-bu'
        assert res.bu_id == 'dept-bu'
        assert res.capacity == 20
        assert res.social_count == 3
        assert res.is_active is True

    def test_import_service_dept_type_needs_no_platform(self, auth_client, bu_hierarchy, _seed_campus_data):
        """客服部门：平台留空即可，编号走 SD 序列。"""
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(
                name='导入客服部', resource_type='客服部门', department='技术BU',
                capacity=40, status='启用',
            ),
        ])
        assert resp.status_code == 200, resp.content
        res = RotationResource.objects.get(name='导入客服部')
        assert res.resource_no == 'SD0001', res.resource_no
        assert res.resource_type == 'SERVICE_DEPT'
        assert res.platform == ''

    def test_import_live_room_requires_platform(self, auth_client, bu_hierarchy, _seed_campus_data):
        """直播间：平台必填（缺平台 → 行级错误）。"""
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='缺平台', resource_type='直播间', department='技术BU'),
        ])
        assert resp.status_code == 400, resp.content
        data = resp.json()['data']
        assert data['skipped'] == 1
        assert any('平台' in e for e in data['errors']), data['errors']

    def test_import_platform_allowed_for_non_live_room(self, auth_client, bu_hierarchy, _seed_campus_data):
        """2026-10-23 闸口1 裁决 Q1：平台对三类形态**均可填**（原口径「非直播间填平台 → 行级错误」作废）。

        客服部门行填中文名平台 → 成功落库且 platform 解析为字典 key。
        """
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='客服带平台', resource_type='客服部门', department='技术BU', platform='视频号'),
        ])
        assert resp.status_code == 200, resp.content
        data = resp.json()['data']
        assert data['skipped'] == 0, data
        assert data['created'] == 1, data
        res = RotationResource.objects.get(name='客服带平台')
        assert res.resource_type == 'SERVICE_DEPT'
        assert res.platform == 'VIDEO_ACCOUNT', '平台中文名应解析为字典 key'

    def test_import_platform_allowed_for_other_type(self, auth_client, bu_hierarchy, _seed_campus_data):
        """闸口1 Q1：OTHER（其他）同样放开平台填写。"""
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='其他带平台', resource_type='其他', department='技术BU', platform='VIDEO_ACCOUNT'),
        ])
        assert resp.status_code == 200, resp.content
        data = resp.json()['data']
        assert data['skipped'] == 0, data
        res = RotationResource.objects.get(name='其他带平台')
        assert res.resource_type == 'OTHER'
        assert res.platform == 'VIDEO_ACCOUNT'

    def test_import_invalid_platform_rejected_for_all_types(self, auth_client, bu_hierarchy, _seed_campus_data):
        """放开类型后，「有值则须为启用中的字典项」对三类形态统一生效。"""
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='平台非法', resource_type='客服部门', department='技术BU', platform='NOT_A_PLATFORM'),
        ])
        assert resp.status_code == 400, resp.content
        data = resp.json()['data']
        assert data['skipped'] == 1, data
        assert any('平台' in e for e in data['errors']), data['errors']
        assert not RotationResource.objects.filter(name='平台非法').exists()

    def test_import_invalid_resource_type_is_skipped(self, auth_client, bu_hierarchy, _seed_campus_data):
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='类型非法', resource_type='办公室', department='技术BU'),
        ])
        assert resp.status_code == 400, resp.content
        data = resp.json()['data']
        assert data['skipped'] == 1
        assert any('资源类型' in e for e in data['errors']), data['errors']

    def test_import_updates_existing_resource_by_no(self, auth_client, bu_hierarchy, _seed_campus_data):
        _reset_resources()
        bu = bu_hierarchy['bu_valid']
        RotationResource.objects.create(
            resource_no='LR0007', name='旧名称', resource_type='LIVE_ROOM',
            department=bu, bu=bu, platform='VIDEO_ACCOUNT',
        )
        resp = _post_import(auth_client, [
            _row(
                resource_no='LR0007', name='新名称', resource_type='直播间',
                department='技术BU', platform='腾讯会议', capacity=50, status='停用',
            ),
        ])
        assert resp.status_code == 200, resp.content
        assert resp.json()['data']['updated'] == 1
        assert resp.json()['data']['created'] == 0

        res = RotationResource.objects.get(resource_no='LR0007')
        assert res.name == '新名称'
        assert res.platform == 'TENCENT_MEETING'
        assert res.capacity == 50
        assert res.is_active is False

    def test_import_type_mismatch_with_existing_is_skipped(self, auth_client, bu_hierarchy, _seed_campus_data):
        """🔴 编号已存在但资源类型不一致 → 跳过（类型决定编号前缀，不允许改）。"""
        _reset_resources()
        bu = bu_hierarchy['bu_valid']
        RotationResource.objects.create(
            resource_no='LR0008', name='直播间', resource_type='LIVE_ROOM',
            department=bu, bu=bu, platform='VIDEO_ACCOUNT',
        )
        resp = _post_import(auth_client, [
            _row(resource_no='LR0008', name='改成客服部', resource_type='客服部门', department='技术BU'),
        ])
        assert resp.status_code == 400, resp.content
        data = resp.json()['data']
        assert data['skipped'] == 1
        assert any('不一致' in e for e in data['errors']), data['errors']
        assert RotationResource.objects.get(resource_no='LR0008').name == '直播间'

    def test_import_unknown_department_is_skipped_with_line_error(self, auth_client, bu_hierarchy, _seed_campus_data):
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='部门不存在', resource_type='直播间', department='不存在的部门', platform='视频号'),
        ])
        assert resp.status_code == 400, resp.content
        body = resp.json()
        assert body['success'] is False
        data = body['data']
        assert data['created'] == 0
        assert data['skipped'] == 1
        assert any('第2行' in e and '归属部门' in e for e in data['errors']), data['errors']
        assert data['errorFile'], '错误行应返回 base64 错误明细文件'
        assert RotationResource.objects.count() == 0

    def test_import_invalid_platform_capacity_status_are_skipped(self, auth_client, bu_hierarchy, _seed_campus_data):
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='平台非法', resource_type='直播间', department='技术BU', platform='不存在平台'),
            _row(name='容量负数', resource_type='直播间', department='技术BU', platform='视频号', capacity=-1),
            _row(name='状态非法', resource_type='直播间', department='技术BU', platform='视频号', status='未知'),
        ])
        assert resp.status_code == 400, resp.content
        data = resp.json()['data']
        assert data['created'] == 0
        assert data['skipped'] == 3, data
        assert len(data['errors']) == 3, data['errors']
        assert RotationResource.objects.count() == 0

    def test_import_ignores_campus_count_column(self, auth_client, bu_hierarchy, _seed_campus_data):
        """校招人数是只读聚合列 → 导入忽略（导出时仍是系统统计值）。"""
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='忽略校招人数', resource_type='直播间', department='技术BU', platform='视频号', campus=999),
        ])
        assert resp.status_code == 200, resp.content
        res = RotationResource.objects.get(name='忽略校招人数')
        assert not hasattr(res, 'campus_count'), '校招人数不应落库'

        rows = _export_rows(auth_client)
        assert len(rows) == 1
        assert rows[0][10] == 0, f'校招人数应为系统统计值（0），实际 {rows[0][10]}'

    def test_import_accepts_free_text_location(self, auth_client, bu_hierarchy, _seed_campus_data):
        """地点为自由文本 → 不校验 LOCATION 字典，原样入库 / 原样导出。"""
        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='自由文本地点', resource_type='直播间', department='技术BU', platform='视频号', location='滨海新区'),
        ])
        assert resp.status_code == 200, resp.content
        assert resp.json()['data']['created'] == 1
        res = RotationResource.objects.get(name='自由文本地点')
        assert res.location == '滨海新区', res.location

        rows = _export_rows(auth_client)
        assert len(rows) == 1
        assert rows[0][7] == '滨海新区', f'导出应原样输出地点，实际 {rows[0][7]!r}'

    def test_import_soft_deleted_resource_is_skipped(self, auth_client, bu_hierarchy, _seed_campus_data):
        _reset_resources()
        bu = bu_hierarchy['bu_valid']
        res = RotationResource.objects.create(
            resource_no='LR0009', name='已删资源', resource_type='LIVE_ROOM',
            department=bu, bu=bu, platform='VIDEO_ACCOUNT',
        )
        res.soft_delete()
        resp = _post_import(auth_client, [
            _row(resource_no='LR0009', name='改名尝试', resource_type='直播间', department='技术BU', platform='视频号'),
        ])
        assert resp.status_code == 400, resp.content
        data = resp.json()['data']
        assert data['skipped'] == 1
        assert any('已删除' in e for e in data['errors']), data['errors']
        assert RotationResource.all_objects.get(resource_no='LR0009').name == '已删资源'

    def test_import_without_file_field_returns_400(self, auth_client):
        resp = auth_client.post(f'{BASE}/resources/import/', {}, format='multipart')
        assert resp.status_code == 400, resp.content
        assert resp.json()['detail'] == '缺少 file 文件字段'

    def test_import_and_export_write_audit_log(self, auth_client, bu_hierarchy, _seed_campus_data):
        """导入 / 导出均写审计（entity=RotationResource）。"""
        from apps.audit.models import AuditLog

        _reset_resources()
        resp = _post_import(auth_client, [
            _row(name='审计资源', resource_type='直播间', department='技术BU', platform='视频号'),
        ])
        assert resp.status_code == 200, resp.content
        assert AuditLog.objects.filter(entity='RotationResource', action='IMPORT').exists()

        auth_client.get(f'{BASE}/resources/export/?file_format=xlsx')
        assert AuditLog.objects.filter(entity='RotationResource', action='EXPORT').exists()

    def test_import_cannot_touch_resource_outside_scope(self, limited_client, bu_hierarchy, _seed_campus_data):
        """行级隔离：scope 外的资源不可 upsert。"""
        _reset_resources()
        bu_out = bu_hierarchy['bu_other']
        RotationResource.objects.create(
            resource_no='LR0021', name='范围外资源', resource_type='LIVE_ROOM',
            department=bu_out, bu=bu_out, platform='VIDEO_ACCOUNT',
        )
        resp = _post_import(limited_client, [
            _row(resource_no='LR0021', name='改名尝试', resource_type='直播间',
                 department='市场BU', platform='视频号'),
        ])
        assert resp.status_code == 400, resp.content
        data = resp.json()['data']
        assert data['skipped'] == 1
        assert any('无权限' in e for e in data['errors']), data['errors']
        assert RotationResource.objects.get(resource_no='LR0021').name == '范围外资源'


# --------------------------------------------------------------------------- #
# 5. 行级隔离（导出）
# --------------------------------------------------------------------------- #
@pytest.mark.django_db
class TestResourceExportScope:
    def test_export_only_contains_resources_in_user_scope(self, limited_client, bu_hierarchy, _seed_campus_data):
        _reset_resources()
        bu_in = bu_hierarchy['bu_valid']
        bu_out = bu_hierarchy['bu_other']
        RotationResource.objects.create(
            name='范围内资源', resource_type='LIVE_ROOM', department=bu_in, bu=bu_in,
            platform='VIDEO_ACCOUNT',
        )
        RotationResource.objects.create(
            name='范围外资源', resource_type='LIVE_ROOM', department=bu_out, bu=bu_out,
            platform='VIDEO_ACCOUNT',
        )
        resp = limited_client.get(f'{BASE}/resources/export/?file_format=xlsx')
        assert resp.status_code == 200, resp.content
        ws = load_workbook(io.BytesIO(resp.content))['轮岗资源']
        names = {r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
        assert '范围内资源' in names
        assert '范围外资源' not in names, '导出不得包含 scope 外的轮岗资源'
